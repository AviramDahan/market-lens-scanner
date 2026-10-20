import hashlib
import io
import json
import shutil
import subprocess
import sys
import zipfile
from contextlib import nullcontext
from pathlib import Path
from types import SimpleNamespace
from urllib.error import HTTPError
from urllib.request import Request

import pytest
from openpyxl import Workbook, load_workbook

from deploy.hetzner import host_archive as archive
from deploy.hetzner import host_archive_runner as runner
from deploy.hetzner.github_archive import ArchiveAPIError, GitHubArchive, SafeRedirect, publish_verified
from deploy.hetzner import archive_checkpoint as checkpoint


@pytest.fixture
def source(tmp_path):
    repo, state, spool = (tmp_path / name for name in ("repo", "state", "spool"))
    (repo / "agent_results/decisions").mkdir(parents=True)
    (repo / "agent_tracker").mkdir()
    state.mkdir()
    (state / "monitor").mkdir()
    (state / "outbox").mkdir()
    (state / "deployed_revision").write_text("source-release")
    (state / "monitor/latest_status.json").write_text('{"run_id":"m1","status":"MONITOR_OK"}')
    (state / "monitor/latest_persisted.json").write_text('{"run_id":"m1","event_count":1}')
    (repo / "agent_results/dashboard_snapshot.json").write_text('{"status":"ok","cash":99000}')
    (repo / "agent_results/decisions/one.jsonl").write_text('{"final_action":"WATCH"}\n')
    (repo / "agent_results/telegram_notifications.jsonl").write_text('{"key":"tp1-existing","sent_at":"2026-10-16T00:00:00Z"}\n')
    wb = Workbook()
    wb.active.title = "Open Positions"
    wb.active.append(["ticker", "qty", "stop", "partial_taken"])
    wb.active.append(["CME", 10, 254.30, True])
    wb.create_sheet("Trade Log").append(["EXIT_STOP", -100])
    wb.save(repo / "agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx")
    return repo, state, spool


def freeze(source, **kwargs):
    return archive.freeze(*source, source_revision="revision", **kwargs)


def manifest(bundle):
    return json.loads((bundle / "manifest.json").read_text())


def local_fetch(spool):
    return lambda generation, name, output: shutil.copyfile(spool / generation / name, output)


class FakeStore:
    repository = "AviramDahan/market-lens-data-archive"

    def __init__(self):
        self.assets = {}
        self.fail_once = False
        self.private = True
        self.corrupt = False

    def require_private(self):
        if not self.private:
            raise ValueError("private required")

    def release(self, tag, create=False):
        return tag.removeprefix("host-v1-")

    def put(self, release, name, path):
        if self.fail_once:
            self.fail_once = False
            raise ArchiveAPIError(502)
        body = path.read_bytes()
        existing = self.assets.get((release, name))
        if existing is not None and existing != body:
            raise ValueError("immutable conflict")
        self.assets[(release, name)] = body

    def get(self, generation, name, output, *, expected_sha=None, expected_bytes=None):
        body = self.assets[(generation, name)]
        if self.corrupt and name.endswith(".zip"):
            body += b"corruption"
        output.write_bytes(body)
        if expected_sha is not None and hashlib.sha256(body).hexdigest() != expected_sha:
            raise ValueError("readback hash failed")
        if expected_bytes is not None and len(body) != expected_bytes:
            raise ValueError("readback size failed")


@pytest.fixture
def checkpoints(source, monkeypatch):
    monkeypatch.setattr(checkpoint, "exclusive", lambda _path: nullcontext())
    monkeypatch.setattr(checkpoint, "require_headroom", lambda *_args: None)
    return checkpoint.Checkpoints(source[0], source[1], store=FakeStore())


def test_checkpoint_bootstrap_and_incremental_proof_are_distinct(source, checkpoints):
    first = checkpoints.save("first", bootstrap=True)
    assert first["proof"]["all_hashes_verified"] is True
    for index in range(20):
        (source[0] / f"agent_results/decisions/history-{index}.jsonl").write_text('{}\n')
    checkpoints.save("second")
    (source[0] / "agent_results/decisions/new.jsonl").write_text('{"action":"WATCH"}\n')
    following = checkpoints.save("third")
    assert following["proof"]["chain_verified"] is True
    assert following["proof"]["all_hashes_verified"] is False
    assert following["proof"]["verified_now_file_count"] < following["proof"]["inventory_file_count"]
    assert checkpoints.recover()["manifest"]["generation"] == following["manifest"]["generation"]
    assert (source[0] / "agent_results/decisions/history-0.jsonl").exists()


def test_failed_checkpoint_cannot_authorize_another_write_and_recovers_exact_generation(source, checkpoints):
    first = checkpoints.save("first", bootstrap=True)
    (source[0] / "agent_results/decisions/new.jsonl").write_text('{}\n')
    checkpoints.store.fail_once = True
    with pytest.raises(ArchiveAPIError):
        checkpoints.save("second")
    pending = json.loads((checkpoints.root / "pending.json").read_text())
    assert checkpoints.receipt()["manifest"] == first["manifest"]
    with pytest.raises(RuntimeError, match="Pending checkpoint"):
        checkpoints.save("must-not-freeze-new")
    receipt = checkpoints.recover()
    assert receipt["manifest"]["generation"] == pending["generation"]
    assert not (checkpoints.root / "pending.json").exists()


def test_bootstrap_upload_failure_is_retryable_without_fake_success(checkpoints):
    checkpoints.store.fail_once = True
    with pytest.raises(ArchiveAPIError):
        checkpoints.save("first", bootstrap=True)
    assert not (checkpoints.root / "last_verified.json").exists()
    assert checkpoints.recover(bootstrap=True)["proof"]["all_hashes_verified"] is True


def test_unbacked_tracker_change_fails_closed_without_rewriting_workbook(source, checkpoints):
    checkpoints.save("first", bootstrap=True)
    tracker = source[0] / checkpoint.TRACKER.removeprefix("repo/")
    tracker.write_bytes(b"UNBACKED PORTFOLIO")
    with pytest.raises(RuntimeError, match="differs from checkpoint"):
        checkpoints.recover()
    assert tracker.read_bytes() == b"UNBACKED PORTFOLIO"


def test_missing_outbox_after_delivery_is_reconciled_before_new_writer(source, checkpoints, monkeypatch):
    from agent import market_lens_ui_agent as ui_agent
    from app import telegram_notifications as telegram
    outbox = source[1] / "outbox/scanner.json"
    body = json.dumps({"schema_version": 1, "records": [
        {"ticker": "TEST", "message": "isolated only", "dedupe_key": "already-sent"}]})
    outbox.write_text(body)
    first = checkpoints.save("first", bootstrap=True)
    assert first["proof"]["delivery_state"] == "VERIFIED_PENDING_DELIVERY"
    outbox.unlink()  # Simulate consumer success immediately before receipt backup.
    receipts = source[0] / checkpoint.RECEIPTS.removeprefix("repo/")
    receipts.write_text(receipts.read_text() + '{"key":"already-sent","sent_at":"2026-10-20T00:00:00Z"}\n')
    recovered = checkpoints.recover()
    assert outbox.read_text() == body
    assert checkpoints.receipts_changed(recovered) is True
    assert "already-sent" in receipts.read_text()  # Recovery never rolls receipts backward.
    monkeypatch.setenv("MARKET_LENS_TELEGRAM_DEDUP_LOG", str(receipts))
    monkeypatch.setenv("MARKET_LENS_TELEGRAM_DEDUP_DAYS", "0")
    monkeypatch.setenv("MARKET_LENS_TELEGRAM_DEDUP_ENABLED", "true")
    settings = telegram.TelegramSettings("isolated-only-token", "isolated-only-destination")
    def notification(text, **kwargs):
        return telegram.send_telegram_notification(text, **kwargs, settings=settings,
            opener=lambda *_args, **_kwargs: pytest.fail("Already-delivered envelope attempted a real send"))
    monkeypatch.setattr(ui_agent, "send_telegram_notification", notification)
    assert [result.status for result in ui_agent.send_buy_notification_outbox(outbox)] == ["duplicate"]
    checkpoints.save("receipt-reconciled")
    outbox.unlink()  # Existing fake consumer checks durable keys; no real send.
    final = checkpoints.save("delivery-complete")
    assert final["proof"]["delivery_state"] == "COMPLETE"
    checkpoints.recover()
    assert not outbox.exists()


def test_unverified_new_outbox_never_overwrites_verified_envelope(source, checkpoints):
    outbox = source[1] / "outbox/scanner.json"
    outbox.write_text('{"schema_version":1,"records":[]}')
    checkpoints.save("first", bootstrap=True)
    outbox.write_text('{"schema_version":1,"records":[],"newer":true}')
    with pytest.raises(RuntimeError, match="outbox is not backed"):
        checkpoints.recover()
    assert json.loads(outbox.read_text())["newer"] is True


def test_corrupt_incremental_asset_keeps_old_verified_receipt(source, checkpoints):
    first = checkpoints.save("first", bootstrap=True)
    (source[0] / "agent_results/decisions/new.jsonl").write_text('{}\n')
    checkpoints.store.corrupt = True
    with pytest.raises(ValueError):
        checkpoints.save("second")
    assert checkpoints.receipt()["manifest"] == first["manifest"]
    checkpoints.store.corrupt = False
    assert checkpoints.recover()["proof"]["chain_verified"] is True


def test_checkpoint_child_timeout_keeps_frozen_generation_for_retry(source, checkpoints, monkeypatch):
    first = checkpoints.save("first", bootstrap=True)
    store = GitHubArchive(FakeStore.repository, "isolated-not-a-real-token")
    monkeypatch.setattr(store, "require_private", lambda: None)
    checkpoints.store = store
    (source[0] / "agent_results/decisions/new.jsonl").write_text('{}\n')
    calls = []
    def timeout(args, **kwargs):
        calls.append((args, kwargs))
        raise subprocess.TimeoutExpired(args, kwargs["timeout"])
    monkeypatch.setattr(checkpoint.subprocess, "run", timeout)
    with pytest.raises(RuntimeError, match="timed out"):
        checkpoints.save("second")
    assert calls[0][1]["timeout"] == 30
    assert "isolated-not-a-real-token" not in str(calls)
    assert checkpoints.receipt()["manifest"] == first["manifest"]
    assert (checkpoints.root / "pending.json").exists()
    assert not list(checkpoints.root.glob("publish-attempt-*"))


def test_incremental_proof_rejects_tampered_verified_base(source, checkpoints):
    from deploy.hetzner.github_archive import incremental_restore_manifest
    previous = checkpoints.save("first", bootstrap=True)
    second = freeze(source, base=previous["manifest"])
    previous["proof"]["manifest_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="verified same-repository"):
        incremental_restore_manifest(manifest(second), previous, checkpoints.store.repository)


def test_upload_only_proof_cannot_roll_authoritative_cursor_backward(source, checkpoints, monkeypatch):
    from deploy.hetzner.github_archive import install_verified_receipt
    first = checkpoints.save("first", bootstrap=True)
    memory_store = checkpoints.store
    real_store = GitHubArchive(FakeStore.repository, "isolated-not-a-real-token")
    monkeypatch.setattr(real_store, "require_private", lambda: None)
    checkpoints.store = real_store
    advanced = []
    def stale_child(args, **kwargs):
        generation = args[args.index("--generation") + 1]
        scratch = Path(args[args.index("--scratch") + 1])
        bundle = checkpoints.root / "generations" / generation
        receipt = publish_verified(bundle, checkpoints.root, memory_store, previous=first, install_cursor=False)
        assert checkpoints.receipt()["manifest"] == first["manifest"]
        (scratch / "verified-result.json").write_text(json.dumps(receipt))
        newer_bundle = archive.freeze(source[0], source[1], checkpoints.root / "generations",
                                      source_revision="newer-owner", base=receipt["manifest"])
        newer = publish_verified(newer_bundle, checkpoints.root, memory_store, previous=receipt, install_cursor=False)
        install_verified_receipt(checkpoints.root, newer)
        advanced.append(newer)
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(checkpoint.subprocess, "run", stale_child)
    with pytest.raises(RuntimeError, match="refusing stale proof"):
        checkpoints.save("second")
    assert checkpoints.receipt()["manifest"] == advanced[0]["manifest"]


def test_roundtrip_preserves_workbook_receipts_outbox_and_cursor(source, tmp_path):
    repo, state, spool = source
    event = {"schema_version": 1, "events": [{"position": {"ticker": "CME"}, "event": {
        "ticker": "CME", "action": "TAKE_PARTIAL_PROFIT", "triggered_at": "2026-10-16T00:00:00Z",
        "trigger_price": 280, "high": 281, "low": 279, "close": 280,
        "quantity": 10, "cash_in": 2800, "note": "existing event"}}]}
    (state / "outbox/monitor.json").write_text(json.dumps(event))
    bundle = freeze(source)
    before = archive.sha256(repo / "agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx")
    restored = tmp_path / "restored"
    proof = archive.restore(manifest(bundle), restored, local_fetch(spool))
    recovered = restored / "repo/agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx"
    assert archive.sha256(recovered) == before
    wb = load_workbook(recovered)
    assert list(wb["Open Positions"].values)[1] == ("CME", 10, 254.3, True)
    wb.close()
    assert json.loads((restored / "runtime/outbox/monitor.json").read_text()) == event
    assert (restored / "runtime/monitor/latest_persisted.json").read_bytes() == (state / "monitor/latest_persisted.json").read_bytes()
    assert (restored / "repo/agent_results/telegram_notifications.jsonl").read_bytes() == (repo / "agent_results/telegram_notifications.jsonl").read_bytes()
    assert archive.sha256(repo / "agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx") == before
    assert proof["all_hashes_verified"] is True
    assert not (restored / ".asset-cache").exists()


def test_incremental_bundle_reuses_verified_file_versions(source, tmp_path):
    repo, _, spool = source
    first = freeze(source)
    base = manifest(first)
    changed = repo / "agent_results/decisions/one.jsonl"
    changed.write_text('{"final_action":"SKIP"}\n')
    second = freeze(source, base=base)
    following = manifest(second)
    current = [entry for entry in following["files"] if entry["generation"] == "self"]
    assert [entry["path"] for entry in current] == ["repo/agent_results/decisions/one.jsonl"]
    inherited = [entry for entry in following["files"] if entry["generation"] != "self"]
    assert all(entry["generation"] == base["generation"] for entry in inherited)
    archive.restore(following, tmp_path / "incremental-restored", local_fetch(spool))
    assert (tmp_path / "incremental-restored/repo/agent_results/decisions/one.jsonl").read_bytes() == changed.read_bytes()
    assert first.exists()  # Old versions remain available.


@pytest.mark.parametrize("ref", ["/app/agent_results/charts/pending.png", "agent_results/charts/pending.png",
                                "/agent-results/charts/pending.png",
                                "https://market-lens.2.28.100.77.sslip.io/agent-results/charts/pending.png"])
def test_media_policy_only_includes_pending_notification_references(source, ref):
    repo, state, _ = source
    (repo / "agent_results/charts").mkdir()
    (repo / "agent_results/charts/unused.png").write_bytes(b"not needed")
    (repo / "agent_results/charts/pending.png").write_bytes(b"pending picture")
    (state / "outbox/scanner.json").write_text(json.dumps({"schema_version": 1,
        "records": [{"ticker": "CME", "message": "existing", "dedupe_key": "existing-key", "chart_ref": ref}]}))
    names = {entry["path"] for entry in manifest(freeze(source))["files"]}
    assert "repo/agent_results/charts/pending.png" in names
    assert "repo/agent_results/charts/unused.png" not in names


def test_credentials_and_host_configuration_are_excluded(source):
    repo, state, _ = source
    (repo / ".env").write_text("password=do-not-copy")
    (state / "host.env").write_text("token=do-not-copy")
    (state / "writer.lock").touch()
    (state / "scanner-staging").mkdir()
    (state / "scanner-staging/temporary.json").write_text("{}")
    names = {entry["path"] for entry in manifest(freeze(source))["files"]}
    assert not any(".env" in name or "staging" in name or "lock" in name for name in names)


def test_credential_like_generated_file_fails_closed(source):
    (source[0] / "agent_results/credentials.json").write_text("{}")
    with pytest.raises(ValueError, match="Credential"):
        freeze(source)


@pytest.mark.parametrize("name", ["../repo/agent_results/a.json", "repo/agent_results/../a.json",
                                  "/repo/agent_results/a.json", "repo\\agent_results\\a.json",
                                  "repo/agent_results/.env", "runtime/secret.json", "repo/agent_results/C:foo.json"])
def test_archive_path_scope_rejects_escape(name):
    with pytest.raises(ValueError):
        archive.safe_name(name)


def test_restore_never_overwrites_existing_directory(source, tmp_path):
    bundle = freeze(source)
    with pytest.raises(ValueError, match="must not exist"):
        archive.restore(manifest(bundle), source[0], lambda *_: pytest.fail("No fetch before destination validation"))


def test_corrupt_download_cannot_become_success_receipt(source, tmp_path):
    store = FakeStore()
    store.corrupt = True
    bundle = freeze(source)
    with pytest.raises(ValueError, match="integrity"):
        publish_verified(bundle, tmp_path / "receipts", store)
    assert not (tmp_path / "receipts/last_verified.json").exists()
    assert not any(name == "verified.json" for _, name in store.assets)


def test_failed_upload_retries_same_generation_without_duplicates(source, tmp_path):
    store = FakeStore()
    store.fail_once = True
    bundle = freeze(source)
    with pytest.raises(ArchiveAPIError):
        publish_verified(bundle, tmp_path / "receipts", store)
    assert not (tmp_path / "receipts/last_verified.json").exists()
    receipt = publish_verified(bundle, tmp_path / "receipts", store)
    uploaded = dict(store.assets)
    duplicate = publish_verified(bundle, tmp_path / "receipts", store)
    assert receipt["manifest"]["generation"] == duplicate["manifest"]["generation"] == bundle.name
    assert store.assets == uploaded
    assert json.loads((tmp_path / "receipts/last_verified.json").read_text())["proof"]["all_hashes_verified"]


def test_public_archive_refused_before_any_upload(source, tmp_path):
    store = FakeStore()
    store.private = False
    with pytest.raises(ValueError, match="private"):
        publish_verified(freeze(source), tmp_path / "receipts", store)
    assert not store.assets
    assert not (tmp_path / "receipts/last_verified.json").exists()


def test_part_splitting_keeps_all_members(source, tmp_path):
    bundle = freeze(source, part_raw_bytes=100)
    assert len(list(bundle.glob("*.zip"))) > 1
    proof = archive.restore(manifest(bundle), tmp_path / "parts-restored", local_fetch(source[2]))
    assert proof["file_count"] == len(manifest(bundle)["files"])


def test_source_mutation_during_freeze_leaves_no_success_bundle(source, monkeypatch):
    original = archive.sha256
    target = source[0] / "agent_results/decisions/one.jsonl"

    def mutate(path):
        digest = original(path)
        if path == target:
            path.write_text("changed during snapshot")
        return digest

    monkeypatch.setattr(archive, "sha256", mutate)
    with pytest.raises(ValueError, match="changed"):
        freeze(source)
    assert list(source[2].iterdir()) == []
    assert target.exists()


def test_missing_execution_marker_is_not_zero_or_valid_backup(source):
    (source[1] / "deployed_revision").unlink()
    with pytest.raises(ValueError, match="Required"):
        freeze(source)


def test_manifest_tamper_rejected(source, tmp_path):
    bundle = freeze(source)
    data = manifest(bundle)
    data["files"][0]["sha256"] = "0" * 64
    with pytest.raises(ValueError, match="identity"):
        archive.restore(data, tmp_path / "tampered", lambda *_: pytest.fail("Do not fetch tampered manifest"))


def test_redirect_strips_auth_for_release_asset_host():
    request = Request("https://api.github.com/repos/owner/repo/releases/assets/1", headers={"Authorization": "Bearer SECRET"})
    redirected = SafeRedirect().redirect_request(request, None, 302, "redirect", {}, "https://release-assets.githubusercontent.com/object?signature=value")
    assert not redirected.has_header("Authorization")
    assert request.has_header("Authorization")
    with pytest.raises(ArchiveAPIError):
        SafeRedirect().redirect_request(request, None, 302, "redirect", {}, "https://attacker.example/object")


def test_api_failure_never_exposes_token_or_signed_url():
    class Fails:
        def open(self, request, **kwargs):
            raise HTTPError("https://example.invalid/?token=SECRET", 403, "SECRET", {}, None)

    store = GitHubArchive("owner/repo", "SECRET", opener=Fails())
    with pytest.raises(ArchiveAPIError) as caught:
        store.json(store.base)
    assert str(caught.value) == "Archive API request failed (status 403)"
    assert "SECRET" not in str(caught.value)


def test_put_leaves_failed_starter_intact_and_uses_retry_name(tmp_path, monkeypatch):
    class Response(io.BytesIO):
        pass

    calls = []

    class Opener:
        def open(self, request, **kwargs):
            calls.append((request.method, request.full_url))
            if request.method == "POST":
                assert request.data.read() == b"payload"
                return Response(b'{}')
            return Response(json.dumps([{"name": "part-0001.zip", "state": "starter", "size": 0}]).encode())

    path = tmp_path / "part.zip"
    path.write_bytes(b"payload")
    GitHubArchive("owner/repo", "dedicated", opener=Opener()).put({"id": 3}, "part-0001.zip", path)
    assert calls[-1][0] == "POST"
    assert "name=part-0001.zip.retry-1" in calls[-1][1]
    assert not any(method == "DELETE" for method, _ in calls)


def test_publisher_lock_is_nonblocking_and_released(monkeypatch, tmp_path):
    events = []

    class FakeFcntl:
        LOCK_EX, LOCK_NB, LOCK_UN = 1, 4, 8

        def flock(self, fd, flags):
            events.append(flags)

    monkeypatch.setattr(runner, "fcntl", FakeFcntl())
    with runner.exclusive(tmp_path / "publisher.lock"):
        assert events == [5]
    assert events == [5, 8]


def test_mirror_failure_is_sanitized_and_does_not_call_trading_runner(monkeypatch, capsys):
    monkeypatch.setattr(runner, "main", lambda: (_ for _ in ()).throw(RuntimeError("sensitive transport error")))
    assert runner.cli() == 1
    output = capsys.readouterr().err
    assert "RuntimeError" in output and "sensitive" not in output


def test_declared_but_missing_portfolio_worksheet_is_not_verified(source):
    workbook = source[0] / "agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx"
    replacement = workbook.with_suffix(".replacement")
    with zipfile.ZipFile(workbook) as original, zipfile.ZipFile(replacement, "w") as target:
        for info in original.infolist():
            if info.filename != "xl/worksheets/sheet1.xml":
                target.writestr(info, original.read(info))
    replacement.replace(workbook)
    with pytest.raises(ValueError, match="worksheet data missing"):
        freeze(source)


@pytest.mark.parametrize("kind,payload", [
    ("monitor", {"schema_version": 999, "events": "corrupt"}),
    ("monitor", {"schema_version": 1, "events": [{"position": {}, "event": {}}]}),
    ("scanner", {"schema_version": 1, "records": [{"ticker": "CME"}]}),
])
def test_unusable_outbox_is_not_a_successful_recovery(source, kind, payload):
    (source[1] / "outbox" / f"{kind}.json").write_text(json.dumps(payload))
    with pytest.raises(ValueError, match="outbox|notification"):
        freeze(source)


def setup_runner(source, monkeypatch):
    repo, state, _ = source
    archive_root = state / "archive"
    archive_root.mkdir()
    monkeypatch.setattr(runner, "REPO", repo)
    monkeypatch.setattr(runner, "STATE", state)
    monkeypatch.setattr(runner, "exclusive", lambda *_: nullcontext())
    monkeypatch.setattr(runner.subprocess, "check_output", lambda *_a, **_k: "revision")
    monkeypatch.setattr(runner.shutil, "disk_usage", lambda *_: SimpleNamespace(free=100 * 1024**3))
    config = {"MARKET_LENS_ARCHIVE_REPOSITORY": "owner/archive", "MARKET_LENS_ARCHIVE_MIN_FREE_GIB": "12"}
    monkeypatch.setattr(runner, "env_value", lambda name: config.get(name, ""))
    return archive_root, config


def test_invalid_live_source_is_retryable_without_creating_pending(source, monkeypatch):
    root, _ = setup_runner(source, monkeypatch)
    dashboard = source[0] / "agent_results/dashboard_snapshot.json"
    dashboard.write_text('{"status":"error"}')
    with pytest.raises(ValueError, match="not healthy"):
        runner.pending_bundle(root)
    assert not (root / "pending.json").exists()
    dashboard.write_text('{"status":"ok"}')
    assert runner.pending_bundle(root).is_dir()


def test_quarantine_retains_failed_bundle_and_allows_new_valid_generation(source, monkeypatch):
    root, _ = setup_runner(source, monkeypatch)
    first = runner.pending_bundle(root)
    next(first.glob("*.zip")).write_bytes(b"damaged local frozen copy")
    with pytest.raises(ValueError, match="integrity"):
        runner.local_restore(first, root / "generations")
    assert runner.pending_bundle(root) == first  # No silent substitution on retry.
    runner.quarantine_pending(root, "operator-review")
    following = runner.pending_bundle(root)
    assert first.exists() and first != following
    assert list((root / "quarantine").glob("*.pending.json"))
    assert runner.local_restore(following, root / "generations")["all_hashes_verified"]


def test_pending_destination_change_fails_before_upload(source, monkeypatch):
    root, config = setup_runner(source, monkeypatch)
    first = runner.pending_bundle(root)
    config["MARKET_LENS_ARCHIVE_REPOSITORY"] = "owner/another-private-repo"
    with pytest.raises(ValueError, match="destination changed"):
        runner.pending_bundle(root)
    assert first.exists()


def test_pending_retry_checks_disk_reserve(source, monkeypatch):
    root, _ = setup_runner(source, monkeypatch)
    runner.pending_bundle(root)
    monkeypatch.setattr(runner.shutil, "disk_usage", lambda *_: SimpleNamespace(free=0))
    with pytest.raises(RuntimeError, match="headroom"):
        runner.pending_bundle(root)


@pytest.mark.parametrize("minimum", ["nan", "inf", "11", "-10"])
def test_invalid_free_space_config_cannot_bypass_reserve(source, monkeypatch, minimum):
    _, config = setup_runner(source, monkeypatch)
    config["MARKET_LENS_ARCHIVE_MIN_FREE_GIB"] = minimum
    with pytest.raises(ValueError, match="at least 12"):
        runner.require_headroom(1)


def test_symlink_source_refused(source):
    link = source[0] / "agent_results/alias.json"
    try:
        link.symlink_to(source[0] / "agent_results/dashboard_snapshot.json")
    except (OSError, NotImplementedError):
        pytest.skip("Symlink creation unavailable")
    with pytest.raises(ValueError, match="Symlink"):
        freeze(source)


@pytest.mark.parametrize("attack", ["traversal", "duplicate", "symlink"])
def test_unsafe_zip_members_are_rejected_even_with_matching_asset_hash(source, tmp_path, attack):
    bundle = freeze(source)
    data = manifest(bundle)
    name = data["files"][0]["asset"]
    path = bundle / name
    first = data["files"][0]["path"]
    with zipfile.ZipFile(path, "a") as target:
        if attack == "symlink":
            info = zipfile.ZipInfo("repo/agent_results/bad.json")
            info.create_system = 3
            info.external_attr = 0o120777 << 16
            target.writestr(info, "target")
        else:
            with pytest.warns(UserWarning, match="Duplicate") if attack == "duplicate" else nullcontext():
                target.writestr("../escape.json" if attack == "traversal" else first, b"bad")
    for entry in data["files"]:
        if entry["asset"] == name:
            entry.update(asset_sha256=archive.sha256(path), asset_bytes=path.stat().st_size)
    del data["generation"]
    data["generation"] = hashlib.sha256(archive.encoded(data)).hexdigest()
    with pytest.raises(ValueError, match="Unsafe|Duplicate"):
        archive.restore(data, tmp_path / "unsafe", lambda generation, asset, output: shutil.copyfile(bundle / asset, output))
    assert not (tmp_path / "escape.json").exists()


def test_missing_inherited_asset_preserves_previous_verified_receipt(source, tmp_path):
    store = FakeStore()
    receipts = tmp_path / "receipts"
    first = freeze(source)
    initial = publish_verified(first, receipts, store)
    saved = (receipts / "last_verified.json").read_bytes()
    (source[0] / "agent_results/decisions/one.jsonl").write_text("{}\n")
    second = freeze(source, base=initial["manifest"])
    asset = next(row["asset"] for row in initial["manifest"]["files"])
    del store.assets[(first.name, asset)]
    with pytest.raises(KeyError):
        publish_verified(second, receipts, store)
    assert (receipts / "last_verified.json").read_bytes() == saved
    assert (second.name, "verified.json") not in store.assets


@pytest.mark.skipif(runner.fcntl is None, reason="Linux flock integration runs in CI/host")
def test_real_publisher_lock_contention_and_exception_release(tmp_path):
    path = tmp_path / "publisher.lock"
    code = "import fcntl,sys; f=open(sys.argv[1],'a'); fcntl.flock(f,fcntl.LOCK_EX|fcntl.LOCK_NB)"
    with pytest.raises(RuntimeError, match="injected"):
        with runner.exclusive(path):
            result = subprocess.run([sys.executable, "-c", code, str(path)], capture_output=True, timeout=10)
            assert result.returncode != 0 and b"BlockingIOError" in result.stderr
            raise RuntimeError("injected")
    assert subprocess.run([sys.executable, "-c", code, str(path)], capture_output=True, timeout=10).returncode == 0


@pytest.mark.parametrize("url", ["http://api.github.com/repos/owner/repo", "https://api.github.com:444/repos/owner/repo",
                                  "https://attacker.example/repos/owner/repo", "https://api.github.com/repos/owner/other"])
def test_token_only_goes_to_https_configured_repo(url):
    class Never:
        def open(self, *_a, **_k):
            pytest.fail("Must reject unsafe URL before attaching token")

    with pytest.raises(ValueError, match="origin"):
        GitHubArchive("owner/repo", "SECRET", opener=Never()).request(url)


def test_restored_scanner_outbox_is_consumable_with_fake_telegram(source, tmp_path, monkeypatch):
    from agent import market_lens_ui_agent as scanner

    record = {"ticker": "CME", "message": "existing saved alert", "dedupe_key": "saved-key"}
    (source[1] / "outbox/scanner.json").write_text(json.dumps({"schema_version": 1, "records": [record]}))
    bundle = freeze(source)
    restored = tmp_path / "consumer-restore"
    archive.restore(manifest(bundle), restored, local_fetch(source[2]))
    seen = []
    monkeypatch.setattr(scanner, "send_buy_notification_records", lambda records: seen.extend(records) or [])
    path = restored / "runtime/outbox/scanner.json"
    before = path.read_bytes()
    assert scanner.send_buy_notification_outbox(path) == []
    assert seen == [record] and path.read_bytes() == before


def test_low_space_and_failed_backup_do_not_touch_active_runner_source():
    root = Path(__file__).resolve().parents[1]
    active = (root / "deploy/hetzner/run_runtime.py").read_text()
    assert "host_archive" not in active
    assert 'push_backup()' in active
    service = (root / "deploy/hetzner/market-lens-archive.service").read_text()
    assert "TimeoutStartSec=1800" in service
    assert "ProtectSystem=strict" in service


def test_numeric_cell_that_openpyxl_cannot_decode_is_not_verified(source):
    workbook = source[0] / "agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx"
    replacement = workbook.with_suffix(".replacement")
    with zipfile.ZipFile(workbook) as original, zipfile.ZipFile(replacement, "w") as target:
        for info in original.infolist():
            body = original.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                assert b"<v>10</v>" in body
                body = body.replace(b"<v>10</v>", b"<v>not-a-number</v>", 1)
            target.writestr(info, body)
    replacement.replace(workbook)
    with pytest.raises(ValueError):
        freeze(source)


def test_incorrect_worksheet_dimensions_cannot_hide_invalid_portfolio_cells(source):
    workbook = source[0] / "agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx"
    replacement = workbook.with_suffix(".replacement")
    with zipfile.ZipFile(workbook) as original, zipfile.ZipFile(replacement, "w") as target:
        for info in original.infolist():
            body = original.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                body = body.replace(b'<dimension ref="A1:D2"', b'<dimension ref="A1:D1"')
                body = body.replace(b'</sheetData>', b'<row r="3"><c r="B3" t="n"><v>not-a-number</v></c></row></sheetData>')
            target.writestr(info, body)
    replacement.replace(workbook)
    with pytest.raises(ValueError):
        freeze(source)


@pytest.mark.parametrize("key,value", [("quantity", "NaN"), ("quantity", True), ("cash_in", float("inf")),
                                       ("trigger_price", "bad"), ("low", float("nan"))])
def test_consumer_breaking_monitor_values_are_not_verified(source, key, value):
    event = {"ticker": "CME", "action": "EXIT_STOP", "triggered_at": "2026-10-16T00:00:00Z",
             "trigger_price": 250.0, "high": 251.0, "low": 249.0, "close": 250.0,
             "quantity": 10, "cash_in": 2500.0, "note": "existing saved event"}
    event[key] = value
    (source[1] / "outbox/monitor.json").write_text(json.dumps({"schema_version": 1,
        "events": [{"position": {"ticker": "CME"}, "event": event}]}))
    with pytest.raises(ValueError, match="notification values"):
        freeze(source)


def test_host_parser_fallback_is_offline_read_only_and_requires_proof(source, monkeypatch):
    from deploy.hetzner import workbook_archive_check as checker

    monkeypatch.setattr(checker, "validate_workbook", lambda *_: (_ for _ in ()).throw(ImportError("not installed")))
    monkeypatch.setattr(archive.os, "getuid", lambda: 1000, raising=False)
    monkeypatch.setattr(archive.os, "getgid", lambda: 1000, raising=False)
    calls = []

    def fake_run(args, **kwargs):
        calls.append((args, kwargs))
        return SimpleNamespace(stdout='{"parser":"openpyxl","validated_sheets":["Open Positions","Trade Log"],"all_worksheets_decoded":true}')

    monkeypatch.setattr(archive.subprocess, "run", fake_run)
    workbook = source[0] / "agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx"
    archive.validate_workbook_parser(workbook)
    args, kwargs = calls[0]
    assert args[args.index("--network") + 1] == "none"
    assert args[args.index("--user") + 1] == "1000:1000"
    assert "--read-only" in args and "--env-file" not in args
    assert kwargs == {"check": True, "text": True, "capture_output": True, "timeout": 120}
    monkeypatch.setattr(archive.subprocess, "run", lambda *_a, **_k: SimpleNamespace(stdout="{}"))
    with pytest.raises(ValueError, match="did not confirm"):
        archive.validate_workbook_parser(workbook)


@pytest.mark.parametrize("key,value", [("quantity", "NaN"), ("quantity", "1.5"), ("entry_price", "NaN"),
                                       ("stop_loss", float("inf")), ("target_1", "bad")])
def test_nested_monitor_position_values_must_be_consumable(source, key, value):
    event = {"ticker": "CME", "action": "TAKE_PARTIAL_PROFIT", "triggered_at": "2026-10-16T00:00:00Z",
             "trigger_price": 280, "high": 281, "low": 279, "close": 280,
             "quantity": 5, "cash_in": 1400, "note": "existing saved event"}
    position = {"ticker": "CME", "quantity": 10, "entry_price": 254.3, "stop_loss": 250}
    position[key] = value
    (source[1] / "outbox/monitor.json").write_text(json.dumps({"schema_version": 1,
        "events": [{"position": position, "event": event}]}))
    with pytest.raises(ValueError, match="position values"):
        freeze(source)
