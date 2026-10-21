import copy
import hashlib
import json
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from deploy.hetzner import host_archive as archive
from deploy.hetzner import local_write_fence as fence
from deploy.hetzner.staged_scanner import prepare_stage


REVISION = "fixture-source-revision"
TRACKER = "repo/agent_tracker/market_lens_agent_portfolio_budget_100k.xlsx"


class FakeCheckpoint:
    """Local manifest/proof fixture; no network, trading worker or Telegram."""

    def __init__(self, repo, state):
        self.repo, self.state = repo, state
        self.root = state / "checkpoints"
        self.root.mkdir()
        self.store = SimpleNamespace(repository="qa/local-write-fence")
        self.calls = []
        self.fail_next_save = False
        self._sequence = 0
        self._pending_receipt = None
        self._install(self.snapshot_receipt())

    def snapshot_receipt(self):
        self._sequence += 1
        bundle = archive.freeze(
            self.repo, self.state, self.root / "fake-remote",
            source_revision=REVISION,
            created_at=f"2026-10-21T00:00:{self._sequence:02d}Z",
        )
        manifest = json.loads((bundle / "manifest.json").read_text())
        proof = {
            "generation": manifest["generation"],
            "manifest_sha256": hashlib.sha256(archive.encoded(manifest)).hexdigest(),
            "repository": self.store.repository,
            "workbook_verified": True,
            "all_hashes_verified": True,
        }
        return {"manifest": manifest, "proof": proof,
                "repository": self.store.repository, "verified_at": "fixture-only"}

    def _install(self, receipt):
        self._verified = copy.deepcopy(receipt)
        (self.root / "last_verified.json").write_bytes(archive.encoded(receipt))

    def receipt(self):
        return copy.deepcopy(self._verified)

    def matching_tracker(self, manifest):
        expected = next(row["sha256"] for row in manifest["files"] if row["path"] == TRACKER)
        if archive.sha256(self.repo / TRACKER.removeprefix("repo/")) != expected:
            raise RuntimeError("Live tracker differs from checkpoint")

    def save(self, revision):
        self.calls.append(("save", revision))
        if self.fail_next_save:
            self.fail_next_save = False
            raise RuntimeError("Fixture freeze failed before pending.json")
        if (self.root / "pending.json").exists():
            raise RuntimeError("Pending fixture checkpoint requires recover")
        receipt = self.snapshot_receipt()
        self._install(receipt)
        return receipt

    def freeze_pending(self):
        self._pending_receipt = self.snapshot_receipt()
        (self.root / "pending.json").write_bytes(archive.encoded({
            "generation": self._pending_receipt["manifest"]["generation"],
            "repository": self.store.repository, "state": "FROZEN",
        }))
        return self._pending_receipt["manifest"]["generation"]

    def recover(self):
        self.calls.append(("recover",))
        if self._pending_receipt is None:
            raise RuntimeError("Fixture has no frozen pending generation")
        self.matching_tracker(self._pending_receipt["manifest"])
        self._install(self._pending_receipt)
        (self.root / "pending.json").unlink()
        return self.receipt()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value), encoding="utf-8")


def live_bytes(checkpoints):
    return {name: path.read_bytes()
            for name, path in archive.recovery_files(checkpoints.repo, checkpoints.state).items()}


def intent(checkpoints):
    return checkpoints.root / "local_write.json"


def assert_recovery_blocked(checkpoints, revision=REVISION):
    before = live_bytes(checkpoints)
    active = intent(checkpoints).read_bytes()
    with pytest.raises((ValueError, RuntimeError)):
        fence.recover_local_write(checkpoints, revision)
    assert checkpoints.calls == []
    assert live_bytes(checkpoints) == before
    assert intent(checkpoints).exists()
    assert intent(checkpoints).read_bytes() == active


@pytest.fixture
def checkpoints(tmp_path, monkeypatch):
    repo, state = tmp_path / "repo", tmp_path / "state"
    (repo / "agent_tracker").mkdir(parents=True)
    (state / "outbox").mkdir(parents=True)
    (state / "monitor").mkdir()
    (state / "deployed_revision").write_text(REVISION)
    write_json(state / "monitor/latest_status.json", {"run_id": "before", "status": "MONITOR_OK"})
    write_json(repo / "agent_results/dashboard_snapshot.json", {"status": "ok", "cash": 99000})
    (repo / "agent_results/telegram_notifications.jsonl").write_text(
        '{"key":"existing-only","sent_at":"2026-10-20T00:00:00Z"}\n'
    )
    workbook = Workbook()
    workbook.active.title = "Open Positions"
    workbook.active.append(["ticker", "quantity", "stop", "partial_taken"])
    workbook.active.append(["FIXTURE", 10, 90, False])
    workbook.create_sheet("Trade Log").append(["fixture-row", 0])
    workbook.save(repo / TRACKER.removeprefix("repo/"))
    workbook.close()
    monkeypatch.setattr(fence, "require_headroom", lambda *_args: None)
    return FakeCheckpoint(repo, state)


@pytest.fixture
def scan(checkpoints, monkeypatch):
    tracker = checkpoints.repo / TRACKER.removeprefix("repo/")
    stage = prepare_stage(checkpoints.repo, checkpoints.state, archive.sha256(tracker))
    workbook = load_workbook(stage.tracker / tracker.name)
    workbook["Open Positions"]["B2"] = 8
    workbook.save(stage.tracker / tracker.name)
    workbook.close()
    dashboard = stage.results / "dashboard_snapshot.json"
    write_json(dashboard, {"status": "ok", "cash": 99200})
    decisions = stage.results / "decisions/new-run.jsonl"
    decisions.write_text('{"ticker":"FIXTURE","final_action":"WATCH"}\n')
    record = stage.results / "runtime/market_lens_agent_fixture.json"
    write_json(record, {"run_id": "fixture", "run_status": "COMPLETE", "result_cards_read": 180})
    outbox = stage.runtime / "outbox/scanner.json"
    write_json(outbox, {"schema_version": 1, "records": []})
    targets = [
        (dashboard, checkpoints.repo / "agent_results/dashboard_snapshot.json"),
        (decisions, checkpoints.repo / "agent_results/decisions/new-run.jsonl"),
        (stage.tracker / tracker.name, tracker),
        (outbox, checkpoints.state / "outbox/scanner.json"),
        (record, checkpoints.repo / "agent_results/runtime" / record.name),
    ]
    monkeypatch.setattr(fence, "promotion_targets", lambda *_args, **_kwargs: targets)
    return SimpleNamespace(stage=stage, targets=targets, tracker=tracker)


def apply_targets(targets):
    for source, target in targets:
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)


def change_workbook(checkpoints):
    path = checkpoints.repo / TRACKER.removeprefix("repo/")
    workbook = load_workbook(path)
    workbook["Open Positions"]["B2"] = 5
    workbook.save(path)
    workbook.close()


def complete_monitor_write(checkpoints):
    change_workbook(checkpoints)
    write_json(checkpoints.state / "monitor/latest_status.json",
               {"run_id": "fixture-after", "status": "MONITOR_OK", "event_count": 1})
    write_json(checkpoints.repo / "agent_results/dashboard_snapshot.json", {"status": "ok", "cash": 99500})
    write_json(checkpoints.state / "outbox/monitor.json", {"schema_version": 1, "events": []})


def test_scanner_low_space_precheck_leaves_live_data_and_no_intent(checkpoints, scan, monkeypatch):
    before = live_bytes(checkpoints)

    def low_space(*_args):
        assert not intent(checkpoints).exists()
        assert live_bytes(checkpoints) == before
        raise RuntimeError("Fixture insufficient headroom")

    monkeypatch.setattr(fence, "require_headroom", low_space)
    with pytest.raises(RuntimeError, match="headroom"):
        fence.begin_scanner(checkpoints, scan.stage, REVISION)
    assert live_bytes(checkpoints) == before
    assert not intent(checkpoints).exists()
    assert checkpoints.calls == []


def test_scanner_intent_exists_before_any_live_file_changes(checkpoints, scan):
    before = live_bytes(checkpoints)
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    assert intent(checkpoints).is_file()
    assert live_bytes(checkpoints) == before
    assert checkpoints.calls == []


def test_recovery_without_intent_is_a_noop(checkpoints):
    before = live_bytes(checkpoints)
    fence.recover_local_write(checkpoints, REVISION)
    assert checkpoints.calls == []
    assert live_bytes(checkpoints) == before


def test_scanner_not_applied_retires_without_checkpoint_or_financial_write(checkpoints, scan):
    before = live_bytes(checkpoints)
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    fence.recover_local_write(checkpoints, REVISION)
    assert not intent(checkpoints).exists()
    assert checkpoints.calls == []
    assert live_bytes(checkpoints) == before


@pytest.mark.parametrize("sealed", [False, True])
def test_complete_scanner_promotion_recovers_exact_live_data(checkpoints, scan, sealed):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    if sealed:
        fence.seal_scanner(checkpoints)
    promoted = live_bytes(checkpoints)
    fence.recover_local_write(checkpoints, REVISION)
    assert checkpoints.calls == [("save", REVISION)]
    assert live_bytes(checkpoints) == promoted
    assert not intent(checkpoints).exists()
    fence.recover_local_write(checkpoints, REVISION)
    assert checkpoints.calls == [("save", REVISION)]


def test_freeze_failure_before_pending_recovers_known_complete_scanner(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    promoted = live_bytes(checkpoints)
    checkpoints.fail_next_save = True
    with pytest.raises(RuntimeError, match="before pending"):
        checkpoints.save(REVISION)
    assert not (checkpoints.root / "pending.json").exists()
    checkpoints.calls.clear()
    fence.recover_local_write(checkpoints, REVISION)
    assert checkpoints.calls == [("save", REVISION)]
    assert live_bytes(checkpoints) == promoted
    assert not intent(checkpoints).exists()


def test_pending_recovery_reuses_frozen_generation_without_new_save(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    generation = checkpoints.freeze_pending()
    promoted = live_bytes(checkpoints)
    fence.recover_local_write(checkpoints, REVISION)
    assert checkpoints.calls == [("recover",)]
    assert checkpoints.receipt()["manifest"]["generation"] == generation
    assert live_bytes(checkpoints) == promoted
    assert not intent(checkpoints).exists()


@pytest.mark.parametrize("target_index", [0, 2, 4])
def test_partial_promotion_is_blocked_without_rollback(checkpoints, scan, target_index):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets([scan.targets[target_index]])
    assert_recovery_blocked(checkpoints)


@pytest.mark.parametrize("target_index", [0, 1, 2, 3, 4])
def test_tampered_promoted_target_cannot_authorize_checkpoint(checkpoints, scan, target_index):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    target = scan.targets[target_index][1]
    if target == scan.tracker:
        change_workbook(checkpoints)
    elif target.suffix == ".jsonl":
        target.write_text('{"unexpected":"tampered"}\n')
    elif target.name == "dashboard_snapshot.json":
        write_json(target, {"status": "ok", "cash": -1})
    elif target.name == "scanner.json":
        write_json(target, {"schema_version": 1, "records": [], "unexpected": True})
    else:
        write_json(target, {"run_status": "COMPLETE", "result_cards_read": 1})
    assert_recovery_blocked(checkpoints)


def test_extra_structured_inventory_blocks_scanner_recovery(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    write_json(checkpoints.repo / "agent_results/unknown.json", {"unknown": True})
    assert_recovery_blocked(checkpoints)


def test_mutated_unchanged_baseline_file_blocks_scanner_recovery(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    (checkpoints.repo / "agent_results/telegram_notifications.jsonl").write_text(
        '{"key":"unexpected","sent_at":"2026-10-21T00:00:00Z"}\n'
    )
    assert_recovery_blocked(checkpoints)


def test_stale_unrelated_base_blocks_complete_promotion(checkpoints, scan):
    unrelated_base = checkpoints.snapshot_receipt()
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    checkpoints._install(unrelated_base)
    assert_recovery_blocked(checkpoints)


def test_different_source_revision_cannot_recover_intent(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    assert_recovery_blocked(checkpoints, revision="unrelated-source")


def test_scanner_seal_refuses_partial_promotion(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets([scan.targets[2]])
    before = live_bytes(checkpoints)
    with pytest.raises((RuntimeError, ValueError)):
        fence.seal_scanner(checkpoints)
    assert live_bytes(checkpoints) == before
    assert checkpoints.calls == []
    assert intent(checkpoints).exists()


def test_finish_requires_verified_checkpoint_not_only_complete_promotion(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    before = live_bytes(checkpoints)
    with pytest.raises((RuntimeError, ValueError)):
        fence.finish_local_write(checkpoints)
    assert live_bytes(checkpoints) == before
    assert checkpoints.calls == []
    assert intent(checkpoints).exists()


def test_finish_rejects_receipt_that_matches_book_but_not_other_targets(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    checkpoints.save(REVISION)
    checkpoints.calls.clear()
    receipt = checkpoints.receipt()
    for entry in receipt["manifest"]["files"]:
        if entry["path"] == "repo/agent_results/dashboard_snapshot.json":
            entry["sha256"] = "0" * 64
    checkpoints._install(receipt)
    before = live_bytes(checkpoints)
    with pytest.raises((RuntimeError, ValueError)):
        fence.finish_local_write(checkpoints)
    assert live_bytes(checkpoints) == before
    assert checkpoints.calls == []
    assert intent(checkpoints).exists()


def test_finish_verified_checkpoint_retires_intent_idempotently(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    checkpoints.save(REVISION)
    before = live_bytes(checkpoints)
    fence.finish_local_write(checkpoints)
    fence.finish_local_write(checkpoints)
    assert live_bytes(checkpoints) == before
    assert not intent(checkpoints).exists()
    assert checkpoints.calls == [("save", REVISION)]


def test_monitor_low_space_blocks_before_worker_changes(checkpoints, monkeypatch):
    before = live_bytes(checkpoints)
    monkeypatch.setattr(fence, "require_headroom",
                        lambda *_args: (_ for _ in ()).throw(RuntimeError("Fixture low space")))
    with pytest.raises(RuntimeError, match="low space"):
        fence.begin_monitor(checkpoints, REVISION)
    assert live_bytes(checkpoints) == before
    assert not intent(checkpoints).exists()
    assert checkpoints.calls == []


def test_monitor_computing_financial_change_fails_closed(checkpoints):
    fence.begin_monitor(checkpoints, REVISION)
    complete_monitor_write(checkpoints)
    assert_recovery_blocked(checkpoints)


@pytest.mark.parametrize("unexpected", ["decisions", "receipts"])
def test_monitor_computing_unknown_structured_change_fails_closed(checkpoints, unexpected):
    fence.begin_monitor(checkpoints, REVISION)
    if unexpected == "decisions":
        path = checkpoints.repo / "agent_results/decisions/unplanned.jsonl"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text('{"unexpected":"not-a-heartbeat"}\n')
    else:
        (checkpoints.repo / "agent_results/telegram_notifications.jsonl").write_text(
            '{"key":"unplanned-receipt","sent_at":"2026-10-21T00:00:00Z"}\n'
        )
    assert_recovery_blocked(checkpoints)


def test_monitor_computing_heartbeat_only_change_is_a_safe_no_event(checkpoints):
    fence.begin_monitor(checkpoints, REVISION)
    write_json(checkpoints.state / "monitor/latest_status.json",
               {"run_id": "new-heartbeat", "status": "MONITOR_OK", "event_count": 0})
    before = live_bytes(checkpoints)
    fence.recover_local_write(checkpoints, REVISION)
    assert live_bytes(checkpoints) == before
    assert not intent(checkpoints).exists()
    assert checkpoints.calls == []


def test_sealed_monitor_recovers_after_freeze_failure_without_reexecution(checkpoints):
    fence.begin_monitor(checkpoints, REVISION)
    complete_monitor_write(checkpoints)
    fence.seal_monitor(checkpoints)
    promoted = live_bytes(checkpoints)
    checkpoints.fail_next_save = True
    with pytest.raises(RuntimeError, match="before pending"):
        checkpoints.save(REVISION)
    checkpoints.calls.clear()
    fence.recover_local_write(checkpoints, REVISION)
    assert checkpoints.calls == [("save", REVISION)]
    assert live_bytes(checkpoints) == promoted
    assert not intent(checkpoints).exists()


def test_sealed_monitor_reuses_existing_pending_checkpoint(checkpoints):
    fence.begin_monitor(checkpoints, REVISION)
    complete_monitor_write(checkpoints)
    fence.seal_monitor(checkpoints)
    generation = checkpoints.freeze_pending()
    before = live_bytes(checkpoints)
    fence.recover_local_write(checkpoints, REVISION)
    assert checkpoints.calls == [("recover",)]
    assert checkpoints.receipt()["manifest"]["generation"] == generation
    assert live_bytes(checkpoints) == before
    assert not intent(checkpoints).exists()


def test_sealed_monitor_extra_inventory_cannot_authorize_checkpoint(checkpoints):
    fence.begin_monitor(checkpoints, REVISION)
    complete_monitor_write(checkpoints)
    fence.seal_monitor(checkpoints)
    write_json(checkpoints.repo / "agent_results/unknown.json", {"unexpected": True})
    assert_recovery_blocked(checkpoints)


def test_no_event_monitor_allows_new_heartbeat_without_checkpoint(checkpoints):
    tracker = checkpoints.repo / TRACKER.removeprefix("repo/")
    before_hash = archive.sha256(tracker)
    before_book = tracker.read_bytes()
    fence.begin_monitor(checkpoints, REVISION)
    write_json(checkpoints.state / "monitor/latest_status.json",
               {"run_id": "new-heartbeat", "status": "MONITOR_OK", "event_count": 0})
    fence.cancel_no_event_monitor(checkpoints, before_hash)
    assert tracker.read_bytes() == before_book
    assert not intent(checkpoints).exists()
    assert checkpoints.calls == []
    assert json.loads((checkpoints.state / "monitor/latest_status.json").read_text())["run_id"] == "new-heartbeat"


@pytest.mark.parametrize("kind", ["scanner", "monitor"])
def test_no_event_monitor_refuses_any_notification_outbox(checkpoints, kind):
    before_hash = archive.sha256(checkpoints.repo / TRACKER.removeprefix("repo/"))
    fence.begin_monitor(checkpoints, REVISION)
    body = {"schema_version": 1, "records": []} if kind == "scanner" else {"schema_version": 1, "events": []}
    write_json(checkpoints.state / f"outbox/{kind}.json", body)
    before = live_bytes(checkpoints)
    with pytest.raises((RuntimeError, ValueError)):
        fence.cancel_no_event_monitor(checkpoints, before_hash)
    assert live_bytes(checkpoints) == before
    assert intent(checkpoints).exists()
    assert checkpoints.calls == []


def test_no_event_monitor_refuses_changed_workbook(checkpoints):
    before_hash = archive.sha256(checkpoints.repo / TRACKER.removeprefix("repo/"))
    fence.begin_monitor(checkpoints, REVISION)
    change_workbook(checkpoints)
    before = live_bytes(checkpoints)
    with pytest.raises((RuntimeError, ValueError)):
        fence.cancel_no_event_monitor(checkpoints, before_hash)
    assert live_bytes(checkpoints) == before
    assert intent(checkpoints).exists()
    assert checkpoints.calls == []


def test_second_begin_cannot_overwrite_unresolved_intent(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    active = intent(checkpoints).read_bytes()
    before = live_bytes(checkpoints)
    with pytest.raises((RuntimeError, ValueError)):
        fence.begin_monitor(checkpoints, REVISION)
    assert intent(checkpoints).read_bytes() == active
    assert live_bytes(checkpoints) == before
    assert checkpoints.calls == []


def reference_old_chart(checkpoints, scan, reference):
    chart = checkpoints.repo / "agent_results/charts/market_lens_agent_old.png"
    chart.parent.mkdir(parents=True, exist_ok=True)
    chart.write_bytes(b"\x89PNG\r\n\x1a\nfixture-old-chart-only")
    write_json(scan.stage.runtime / "outbox/scanner.json", {"schema_version": 1, "records": [
        {"ticker": "FIXTURE", "message": "fixture-only", "dedupe_key": "fixture-old-chart",
         "chart_path": reference},
    ]})
    assert all(target != chart for _, target in scan.targets)
    return chart


@pytest.mark.parametrize("reference", [
    "/agent-results/charts/market_lens_agent_old.png",
    "https://market-lens.2.28.100.77.sslip.io/agent-results/charts/market_lens_agent_old.png",
    "/app/agent_results/charts/market_lens_agent_old.png",
])
def test_scanner_outbox_can_reuse_declared_old_chart_not_promoted_this_run(checkpoints, scan, reference):
    chart = reference_old_chart(checkpoints, scan, reference)
    original = chart.read_bytes()
    name = "repo/agent_results/charts/market_lens_agent_old.png"
    assert name not in archive.recovery_files(checkpoints.repo, checkpoints.state)
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    declared = json.loads(intent(checkpoints).read_text())
    assert declared["dependencies"][name] == archive.sha256(chart)
    assert name not in declared["before_files"]
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    before = live_bytes(checkpoints)
    fence.recover_local_write(checkpoints, REVISION)
    assert checkpoints.calls == [("save", REVISION)]
    assert live_bytes(checkpoints) == before
    assert chart.read_bytes() == original
    verified = {row["path"]: row["sha256"] for row in checkpoints.receipt()["manifest"]["files"]}
    assert verified[name] == archive.sha256(chart)
    assert not intent(checkpoints).exists()


def test_missing_old_chart_dependency_blocks_begin_before_live_changes(checkpoints, scan):
    chart = reference_old_chart(checkpoints, scan, "/agent-results/charts/market_lens_agent_old.png")
    chart.unlink()
    before = live_bytes(checkpoints)
    with pytest.raises((RuntimeError, ValueError)):
        fence.begin_scanner(checkpoints, scan.stage, REVISION)
    assert live_bytes(checkpoints) == before
    assert checkpoints.calls == []
    assert not intent(checkpoints).exists()


def test_changed_declared_old_chart_blocks_complete_scanner_recovery(checkpoints, scan):
    chart = reference_old_chart(checkpoints, scan, "/agent-results/charts/market_lens_agent_old.png")
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    chart.write_bytes(b"\x89PNG\r\n\x1a\nchanged-old-chart")
    assert_recovery_blocked(checkpoints)


def test_old_chart_dependency_does_not_allow_extra_structured_inventory(checkpoints, scan):
    reference_old_chart(checkpoints, scan, "/agent-results/charts/market_lens_agent_old.png")
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    write_json(checkpoints.repo / "agent_results/unplanned.json", {"unexpected": True})
    assert_recovery_blocked(checkpoints)


@pytest.mark.parametrize("corruption", ["before_tracker_identity", "duplicate_target"])
def test_corrupt_intent_identity_or_duplicate_target_is_rejected(checkpoints, scan, corruption):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    value = json.loads(intent(checkpoints).read_text())
    if corruption == "before_tracker_identity":
        value["before_tracker_sha256"] = "0" * 64
    else:
        value["files"].append(dict(value["files"][0]))
    write_json(intent(checkpoints), value)
    assert_recovery_blocked(checkpoints)


def test_no_event_history_is_compact_and_binds_original_inventory(checkpoints):
    before_hash = archive.sha256(checkpoints.repo / TRACKER.removeprefix("repo/"))
    fence.begin_monitor(checkpoints, REVISION)
    value = json.loads(intent(checkpoints).read_text())
    fence.cancel_no_event_monitor(checkpoints, before_hash)
    history = checkpoints.state / "coordination/write-fence-history" / (value["id"] + ".json")
    record = json.loads(history.read_text())
    assert record["resolution"] == "NO_FINANCIAL_WRITE"
    assert record["inventory_count"] == len(value["before_files"])
    assert record["inventory_sha256"] == hashlib.sha256(archive.encoded(value["before_files"])).hexdigest()
    assert record["intent_sha256"] == hashlib.sha256(archive.encoded(value)).hexdigest()
    assert record["before_tracker_sha256"] == record["after_tracker_sha256"] == before_hash
    assert "before_files" not in record
    assert history.stat().st_size < 2048


def test_history_write_failure_preserves_active_intent(checkpoints, monkeypatch):
    before_hash = archive.sha256(checkpoints.repo / TRACKER.removeprefix("repo/"))
    fence.begin_monitor(checkpoints, REVISION)
    original = intent(checkpoints).read_bytes()
    def fail_history(*_args):
        raise OSError("fixture history storage failure")
    monkeypatch.setattr(fence, "_store", fail_history)
    with pytest.raises(OSError, match="history storage"):
        fence.cancel_no_event_monitor(checkpoints, before_hash)
    assert intent(checkpoints).read_bytes() == original


def test_financial_history_keeps_full_verified_checkpoint_evidence(checkpoints, scan):
    fence.begin_scanner(checkpoints, scan.stage, REVISION)
    apply_targets(scan.targets)
    fence.seal_scanner(checkpoints)
    value = json.loads(intent(checkpoints).read_text())
    fence.recover_local_write(checkpoints, REVISION)
    history = checkpoints.state / "coordination/write-fence-history" / (value["id"] + ".json")
    record = json.loads(history.read_text())
    receipt = checkpoints.receipt()
    assert record["before_files"] == value["before_files"]
    assert record["files"] == value["files"]
    assert record["verified_generation"] == receipt["manifest"]["generation"]
    assert record["verified_manifest_sha256"] == receipt["proof"]["manifest_sha256"]
