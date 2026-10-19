import json
from datetime import datetime, timedelta

from openpyxl import Workbook

from agent import market_lens_ui_agent as agent
from app import agent_dashboard


def workbook_with_trade_log() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trade Log"
    sheet.append([f"column_{index}" for index in range(32)])
    return workbook


def test_single_pass_followup_projection_matches_all_legacy_queues(tmp_path, monkeypatch):
    wb = Workbook()
    wb.active.title = "Open Positions"
    wb.active.append(["Ticker"])
    wb.active.append(["OPEN"])
    ws = wb.create_sheet("Setup Watchlist")
    ws.append([f"column_{i}" for i in range(18)])
    now = datetime.now().replace(microsecond=0)
    for i, (ticker, action, data) in enumerate([
        ("READY", "WATCH", {"watch_status": "WATCH_READY"}),
        ("WATCH", "WATCH", {"reason": "Needs confirmation"}),
        ("NEAR", "SKIP", {"setup_type": "VWAP", "setup_score": .4, "net_rr": 1.7}),
        ("OLD", "WATCH", {}),
        ("UPDATED", "WATCH", {"watch_status": "WATCH_READY"}),
        ("UPDATED", "SKIP", {"setup_type": "No Trade"}),
        ("WARNING_READY", "WATCH", {"warnings": ["WATCH_READY: staged"]}),
    ]):
        row = [None] * 18
        row[0] = (now - timedelta(minutes=10-i) if ticker != "OLD" else now-timedelta(days=30)).isoformat()
        row[1], row[2], row[12] = ticker, "No Trade" if ticker == "UPDATED" and action == "SKIP" else "VWAP", action
        row[17] = json.dumps({**data, "unused_large_evidence": "x" * 10000})
        ws.append(row)
    path = tmp_path / "followups.xlsx"
    wb.save(path)
    before = path.read_bytes()
    expected = {"open": agent.read_open_position_tickers(path), "ready": agent.read_recent_watch_ready_tickers(path),
                "watch": agent.read_recent_watch_tickers(path), "near": agent.read_recent_near_miss_tickers(path),
                "skip": agent.read_recent_skip_tickers(path), "today": agent.read_today_scanned_tickers(path)}
    actual_load = agent.load_workbook
    calls = []
    def counted_load(*args, **kwargs):
        calls.append(1)
        return actual_load(*args, **kwargs)
    monkeypatch.setattr(agent, "load_workbook", counted_load)
    assert agent.read_scan_followup_tickers(path) == expected
    assert len(calls) == 1
    assert path.read_bytes() == before


def test_missing_followup_workbook_fails_visibly_instead_of_empty_scan(tmp_path):
    import pytest
    with pytest.raises(FileNotFoundError):
        agent.read_scan_followup_tickers(tmp_path / "missing.xlsx")


def append_trade(
    workbook: Workbook,
    *,
    timestamp: datetime,
    action: str,
    ticker: str,
    quantity: int,
    cash_out: float = 0.0,
    cash_in: float = 0.0,
    decision: dict | None = None,
) -> None:
    row = [None] * 32
    row[0] = timestamp.isoformat(timespec="seconds")
    row[1] = action
    row[2] = ticker
    row[3] = 100.0
    row[4] = 110.0 if action != "BUY_SIMULATED" else None
    row[5] = quantity
    row[6] = 1.0
    row[7] = cash_out
    row[8] = cash_in
    row[9] = cash_out
    row[10] = cash_in
    row[11] = 95.0
    row[12] = 110.0
    row[13] = 120.0
    row[14] = 50.0
    row[19] = json.dumps(decision or {})
    row[20] = f"{ticker}-{timestamp:%Y%m%d%H%M%S}"
    workbook["Trade Log"].append(row)


def test_trade_analytics_reader_matches_dashboard_reader_without_universe_refresh(monkeypatch) -> None:
    workbook = workbook_with_trade_log()
    now = datetime.now().replace(microsecond=0)
    append_trade(
        workbook,
        timestamp=now - timedelta(days=2),
        action="BUY_SIMULATED",
        ticker="ABC",
        quantity=10,
        cash_out=1_000.0,
        decision={"setup_type": "Fib", "sector": "Technology"},
    )
    append_trade(
        workbook,
        timestamp=now - timedelta(days=1),
        action="TAKE_PROFIT",
        ticker="ABC",
        quantity=10,
        cash_in=1_100.0,
        decision={"setup_type": "Fib", "sector": "Technology"},
    )
    monkeypatch.setattr(agent_dashboard, "_SECTOR_MAP", {"ABC": "Technology"})

    expected = agent_dashboard.read_trades(workbook)
    actual = agent_dashboard.read_trade_analytics(
        workbook,
        ticker_sectors={"ABC": "Technology"},
    )

    assert actual == expected
    assert agent_dashboard.compute_cash(actual, 100_000.0) == agent.compute_cash(workbook, 100_000.0)
    assert agent_dashboard.compute_realized_pnl(actual) == agent_dashboard.compute_realized_pnl(expected)
    assert agent_dashboard.compute_full_trade_performance(actual) == (
        agent_dashboard.compute_full_trade_performance(expected)
    )


def test_trade_ledger_state_matches_legacy_cash_cooldown_and_pilot_counts() -> None:
    workbook = workbook_with_trade_log()
    now = datetime.now().replace(microsecond=0)
    append_trade(
        workbook,
        timestamp=now,
        action="BUY_SIMULATED",
        ticker="PILOT",
        quantity=10,
        cash_out=1_000.0,
        decision={"entry_mode": "neutral_pilot", "setup_type": "VWAP"},
    )
    append_trade(
        workbook,
        timestamp=now - timedelta(days=1),
        action="EXIT_STOP",
        ticker="STOPPED",
        quantity=5,
        cash_in=475.0,
        decision={"setup_type": "Breakout"},
    )

    state = agent.read_trade_ledger_state(
        workbook,
        starting_capital=100_000.0,
        cooldown_days=3,
        now=now,
    )

    assert state["cash"] == agent.compute_cash(workbook, 100_000.0)
    assert state["neutral_pilot_buys_today"] == agent.count_neutral_pilot_buys_today(workbook)
    assert state["recent_stop_events"] == agent.read_recent_stop_events(workbook, 3)
