from copy import deepcopy
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app import strategy
from agent import market_lens_ui_agent as agent


def position(quantity=10):
    return dict(ticker="OLD", quantity=quantity, entry_price=100., current_price=100.,
                stop_loss=95., target_1=110., target_2=120., exposure_ils=quantity * 100.,
                risk_ils=quantity * 5., partial_taken=False)


def candidate(ticker="OLD", price=120.):
    return agent.SetupResult(ticker=ticker, setup_type="Fib", score=.6, current_price=price,
                           buy_zone_low=99., buy_zone_high=101., stop_loss=95., target_1=110.,
                           target_2=120., risk_reward=3., reason="fixture", raw_text="")


@pytest.mark.parametrize("apply", [strategy.apply_strategy_exit, agent.apply_exit_decision])
@pytest.mark.parametrize("quantity", [1, 5, 10])
def test_partial_accounting_and_duplicate_protection(apply, quantity):
    positions = {"OLD": position(quantity)}
    closed = max(1, quantity // 2)
    decision = strategy.StrategyDecision("TAKE_PARTIAL_PROFIT", "TP1", quantity=closed,
                                         cash_in_ils=closed * 110., execution_price=110.)
    cash_delta, exposure_delta = apply(positions, candidate(price=110.), decision, 1.)
    assert cash_delta == closed * 110
    assert exposure_delta == (quantity - closed) * 110 - quantity * 100
    if quantity > 1:
        assert positions["OLD"]["quantity"] == quantity - closed
        assert positions["OLD"]["stop_loss"] == 100
        assert positions["OLD"]["partial_taken"] is True
    else:
        assert not positions
    before = deepcopy(positions)
    with pytest.raises(ValueError):
        apply(positions, candidate(price=110.), decision, 1.)
    assert positions == before


@pytest.mark.parametrize("apply", [strategy.apply_strategy_exit, agent.apply_exit_decision])
def test_full_exit_releases_entire_marked_exposure(apply):
    positions = {"OLD": position()}
    decision = strategy.StrategyDecision("TAKE_PROFIT", "TP2", quantity=10,
                                         cash_in_ils=1200., execution_price=120.)
    assert apply(positions, candidate(), decision, 1.) == (1200., -1000.)
    assert not positions


@pytest.mark.parametrize("quantity", [0, -1, 11])
def test_invalid_exit_quantity_is_atomic(quantity):
    positions = {"OLD": position()}
    before = deepcopy(positions)
    decision = strategy.StrategyDecision("EXIT_STOP", "fixture", quantity=quantity, cash_in_ils=950.)
    with pytest.raises(ValueError):
        strategy.apply_strategy_exit(positions, candidate(price=95.), decision, 1.)
    assert positions == before


def test_stop_proceeds_and_exposure_use_portfolio_currency():
    positions = {"OLD": position()}
    positions["OLD"]["exposure_ils"] = 3700.
    decision = strategy.StrategyDecision("EXIT_STOP", "stop", quantity=10,
                                         cash_in_ils=3515., execution_price=95.)
    assert strategy.apply_strategy_exit(positions, candidate(price=95.), decision, 3.7) == (3515., -3700.)


def test_in_memory_proceeds_match_ledger_without_double_credit(tmp_path):
    wb = Workbook()
    wb.active.title = "Trade Log"
    wb.active.append(["header"] * 32)
    buy = strategy.StrategyDecision("BUY_SIMULATED", "entry", quantity=10, cash_out_ils=1000.)
    agent.append_trade_log_row(wb, "2026-09-08T14:00:00Z", candidate(price=100.), buy, 1., tmp_path / "unused.png")
    cash = agent.compute_cash(wb, 100_000.)
    assert cash == 99_000.
    holdings = {"OLD": position()}
    exit_decision = strategy.StrategyDecision("TAKE_PROFIT", "TP2", quantity=10,
                                              cash_in_ils=1200., execution_price=120.)
    agent.capture_position_exit_plan(exit_decision.decision_json, holdings["OLD"])
    cash_delta, _ = agent.apply_exit_decision(holdings, candidate(), exit_decision, 1.)
    assert wb["Trade Log"].max_row == 2
    agent.append_trade_log_row(wb, "2026-09-09T14:00:00Z", candidate(), exit_decision, 1., tmp_path / "unused.png")
    assert agent.compute_cash(wb, 100_000.) == cash + cash_delta == 100_200.
    assert wb["Trade Log"].max_row == 3


def context():
    return SimpleNamespace(sector_health={"Technology": {}},
                           config=SimpleNamespace(stop_cooldown_days=3),
                           market_regime=SimpleNamespace(max_total_exposure=40_000.))


def test_hold_mark_delta_is_idempotent_and_uses_portfolio_currency():
    p = position()
    p["exposure_ils"] = 3700.
    assert strategy.refresh_strategy_position(p, 105., 3.7) == 185.
    assert p["risk_ils"] == 370.
    assert p["unrealized_ils"] == 185.
    assert strategy.refresh_strategy_position(p, 105., 3.7) == 0.


@pytest.mark.parametrize("price", [0., float("nan"), float("inf")])
def test_invalid_hold_price_does_not_mutate_position(price):
    p = position()
    before = deepcopy(p)
    with pytest.raises(ValueError):
        strategy.refresh_strategy_position(p, price, 1.)
    assert p == before


def test_user_processes_existing_first_but_preserves_display_order(monkeypatch):
    monkeypatch.setattr(strategy, "build_agent_run_context", lambda **_: context())
    monkeypatch.setattr(strategy, "base_universe", lambda: {})
    observed = []
    def final(**kwargs):
        observed.append(kwargs["result"].ticker)
        if kwargs["result"].ticker == "NEW":
            assert kwargs["cash_available"] == 1200.
        return {"final_action": kwargs["initial_action"], "reason": "fixture"}
    monkeypatch.setattr(strategy, "evaluate_agent_candidate", final)
    rows = [candidate("NEW", 100.), candidate("OLD", 120.)]
    output = strategy.apply_strategy_decisions(rows, analysis_period="6mo", min_rr=2,
                                               open_positions={"OLD": position()}, cash=0., exposure=1000.)
    assert observed == ["OLD", "NEW"]
    assert [r.ticker for r in output] == ["NEW", "OLD"]


@pytest.mark.parametrize("price,cash,exposure,quantity", [(120., 1200., 0., 12), (110., 550., 550., 5), (105., 0., 1050., 0)])
def test_user_next_candidate_sees_exit_proceeds_without_mutating_input(monkeypatch, price, cash, exposure, quantity):
    original = {"OLD": position()}
    before = deepcopy(original)
    observed = []
    monkeypatch.setattr(strategy, "build_agent_run_context", lambda **_: context())
    monkeypatch.setattr(strategy, "base_universe", lambda: {})

    def final(**kwargs):
        if kwargs["result"].ticker == "NEW":
            if price == 105:
                assert kwargs["portfolio_open_risk_before"] == 100.
            observed.append((kwargs["cash_available"], kwargs["portfolio_exposure_before"], kwargs["quantity"]))
            return {"final_action": "WATCH", "reason": "Final gate remains authoritative"}
        return {"final_action": kwargs["initial_action"], "reason": "Exit"}

    monkeypatch.setattr(strategy, "evaluate_agent_candidate", final)
    strategy.apply_strategy_decisions([candidate(price=price), candidate("NEW", 100.)],
                                      analysis_period="6mo", min_rr=2, open_positions=original,
                                      cash=0., exposure=1000.)
    assert observed == [(cash, exposure, quantity)]
    assert original == before


@pytest.mark.parametrize("price,cash,exposure,quantity", [(120., 1200., 0., 12), (110., 550., 550., 5), (105., 0., 1050., 0)])
def test_agent_next_candidate_sees_exit_proceeds(monkeypatch, tmp_path, price, cash, exposure, quantity):
    settings = SimpleNamespace(excel_path=tmp_path / "unused.xlsx", min_rr=2,
                               analysis_period="6mo", universe="smart", tickers=[])
    monkeypatch.setattr(agent, "load_workbook", lambda _: object())
    monkeypatch.setattr(agent, "ensure_agent_columns", lambda _: None)
    monkeypatch.setattr(agent, "read_settings", lambda _: {})
    monkeypatch.setattr(agent, "build_agent_run_context", lambda **_: context())
    monkeypatch.setattr(agent, "base_universe", lambda: {})
    monkeypatch.setattr(agent, "read_recent_stop_events", lambda *_: {})
    monkeypatch.setattr(agent, "count_neutral_pilot_buys_today", lambda _: 0)
    monkeypatch.setattr(agent, "read_open_positions", lambda _: {"OLD": position()})
    monkeypatch.setattr(agent, "compute_cash", lambda *_: 0.)
    monkeypatch.setattr(agent, "rank_results_for_allocation", lambda rows, _: rows)
    monkeypatch.setattr(agent, "calculate_setup_score_percentiles", lambda _: {})
    monkeypatch.setattr(agent, "evaluate_shadow_strategies", lambda *_: [])
    monkeypatch.setattr(agent, "build_selection_context", lambda *args, **kwargs: "")
    observed = []

    class EndOfCheck(Exception):
        pass

    def final(**kwargs):
        if kwargs["result"].ticker == "NEW":
            if price == 105:
                assert kwargs["portfolio_open_risk_before"] == 100.
            observed.append((kwargs["cash_available"], kwargs["portfolio_exposure_before"], kwargs["quantity"]))
            raise EndOfCheck
        return {"final_action": kwargs["initial_action"], "reason": "Exit"}

    monkeypatch.setattr(agent, "evaluate_agent_candidate", final)
    with pytest.raises(EndOfCheck):
        agent.update_workbook(settings=settings, run_id="fixture", results=[candidate(price=price), candidate("NEW", 100.)],
                              screenshot_path=tmp_path / "unused.png", summary_path=tmp_path / "unused.md",
                              decision_path=tmp_path / "unused.jsonl", errors=[])
    assert observed == [(cash, exposure, quantity)]
    assert not list(tmp_path.iterdir())
