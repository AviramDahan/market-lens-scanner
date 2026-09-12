import json
from types import SimpleNamespace
from datetime import datetime

import pandas as pd
import pytest
from openpyxl import Workbook

from agent.position_monitor import (
    MonitorSettings, apply_event, build_event, ensure_position_events_sheet,
    monitor_position, position_trade_id, read_last_event_times,
    process_position_window, classify_monitor_results,
)
from agent.market_lens_ui_agent import capture_position_exit_plan, append_trade_log_row


def test_scanner_exit_ledger_preserves_original_plan(tmp_path):
    p = position()
    payload = {}
    capture_position_exit_plan(payload, p)
    p["stop_loss"] = p["entry_price"]
    result = SimpleNamespace(ticker="CHTR", current_price=160, stop_loss=155,
                             target_1=170, target_2=190, chart_url="", selection_context="")
    decision = SimpleNamespace(action="TAKE_PARTIAL_PROFIT", execution_price=159.44,
                               quantity=13, cash_in_ils=2072.72, cash_out_ils=0,
                               risk_ils=0, feedback="TP1", decision_json=payload)
    wb = book()
    append_trade_log_row(wb, "2026-09-08T14:00:00Z", result, decision, 1, tmp_path / "chart.png")
    ws = wb["Trade Log"]
    assert [ws.cell(2, c).value for c in (12, 13, 14, 21)] == [146.49, 159.44, 180.62, "new-trade"]


def position():
    return dict(ticker="CHTR", entry_date="2026-09-04T14:06:53Z", entry_price=150.36,
                quantity=26, stop_loss=146.49, target_1=159.44, target_2=180.62,
                decision_json=json.dumps({"trade_id": "new-trade"}), partial_taken=False)


def book():
    wb = Workbook()
    wb.active.title = "Trade Log"
    wb.active.append(["Header"] * 32)
    ensure_position_events_sheet(wb)
    return wb


def settings(tmp_path):
    return MonitorSettings(tmp_path / "unused.xlsx", tmp_path, "5d", "1m", False, "")


@pytest.mark.parametrize("terminal,high,low,expected_price", [
    ("EXIT_STOP", 156, 149, 150.36),
    ("TAKE_PROFIT", 182, 175, 180.62),
])
def test_window_applies_partial_then_later_terminal_once(monkeypatch, tmp_path, terminal, high, low, expected_price):
    frame = pd.DataFrame({"High": [160, high], "Low": [155, low], "Close": [159, high - 1]},
                         index=pd.to_datetime(["2026-09-08T14:11:00Z", "2026-09-08T14:12:00Z"]))
    calls = []
    def fetch(*args, **kwargs):
        calls.append(1)
        return frame
    monkeypatch.setattr("agent.position_monitor.fetch_intraday_frame", fetch)
    wb, p = book(), position()
    positions = {"CHTR": p}
    results, notifications = process_position_window(
        wb, positions, p, settings=settings(tmp_path), since=None,
        currency_rate=1, timestamp="2026-09-08T14:15:00Z", run_id="window",
    )
    assert [r.event.action for r in results] == ["TAKE_PARTIAL_PROFIT", terminal]
    assert [r.event.quantity for r in results] == [13, 13]
    assert results[-1].event.trigger_price == expected_price
    assert sum(r.event.cash_in for r in results) == pytest.approx(round(13 * (159.44 + expected_price), 2))
    assert len(calls) == 1
    assert positions == {}
    assert wb["Trade Log"].max_row == 3
    assert wb["Position Events"].max_row == 3
    assert notifications[0][0]["quantity"] == 26
    assert notifications[1][0]["quantity"] == 13
    assert notifications[1][0]["stop_loss"] == 150.36
    assert classify_monitor_results(results)["positions_checked"] == 1


def test_window_partial_then_hold_refreshes_and_replay_is_noop(monkeypatch, tmp_path):
    frame = pd.DataFrame({"High": [160, 165], "Low": [155, 158], "Close": [159, 164]},
                         index=pd.to_datetime(["2026-09-08T14:11:00Z", "2026-09-08T14:12:00Z"]))
    monkeypatch.setattr("agent.position_monitor.fetch_intraday_frame", lambda *a, **k: frame)
    wb, p = book(), position()
    positions = {"CHTR": p}
    args = dict(settings=settings(tmp_path), currency_rate=1,
                timestamp="2026-09-08T14:15:00Z", run_id="window")
    results, _ = process_position_window(wb, positions, p, since=None, **args)
    assert len(results) == 1
    assert p["current_price"] == 164
    assert p["quantity"] == 13
    cursor = read_last_event_times(wb, positions)["CHTR"]
    again, notifications = process_position_window(wb, positions, p, since=cursor, **args)
    assert again[0].status == "HOLD"
    assert not notifications
    assert wb["Trade Log"].max_row == 2


def test_window_does_not_infer_new_stop_order_inside_tp1_candle(monkeypatch, tmp_path):
    frame = pd.DataFrame({"High": [160], "Low": [149], "Close": [159]},
                         index=pd.to_datetime(["2026-09-08T14:11:00Z"]))
    monkeypatch.setattr("agent.position_monitor.fetch_intraday_frame", lambda *a, **k: frame)
    wb, p = book(), position()
    results, _ = process_position_window(
        wb, {"CHTR": p}, p, settings=settings(tmp_path), since=None,
        currency_rate=1, timestamp="2026-09-08T14:15:00Z", run_id="window",
    )
    assert len(results) == 1
    assert results[0].event.action == "TAKE_PARTIAL_PROFIT"
    assert p["quantity"] == 13


def test_previous_trade_stop_bar_cannot_close_reentry(monkeypatch, tmp_path):
    frame = pd.DataFrame({"High": [146.94, 151, 152], "Low": [146.365, 150, 151],
                          "Close": [146.365, 150.5, 151.5]},
                         index=pd.to_datetime(["2026-09-01T14:11:00Z", "2026-09-04T14:07:00Z", "2026-09-08T14:00:00Z"]))
    monkeypatch.setattr("agent.position_monitor.fetch_intraday_frame", lambda *a, **k: frame)
    p = position()
    result = monitor_position(p, settings=settings(tmp_path),
                              since=datetime.fromisoformat("2026-09-01T13:30:00+00:00"), currency_rate=1)
    assert result.status == "HOLD" and result.event is None
    assert json.loads(p["decision_json"])["mae"] < 26


def test_missing_entry_is_not_replaced_with_old_event_cursor(tmp_path):
    p = position()
    p["entry_date"] = ""
    result = monitor_position(p, settings=settings(tmp_path), since=None, currency_rate=1)
    assert result.status == "DATA_ERROR"


@pytest.mark.parametrize("high,low,close", [
    (float("nan"), 150, 151), (152, 150, float("inf")),
    (149, 151, 150), (152, 150, 153), (152, 0, 151),
])
def test_bad_price_bar_is_data_error_not_hold_or_exit(monkeypatch, tmp_path, high, low, close):
    frame = pd.DataFrame({"High": [high], "Low": [low], "Close": [close]},
                         index=pd.to_datetime(["2026-09-08T14:11:00Z"]))
    monkeypatch.setattr("agent.position_monitor.fetch_intraday_frame", lambda *a, **k: frame)
    p = position()
    before = dict(p)
    result = monitor_position(p, settings=settings(tmp_path), since=None, currency_rate=1)
    assert result.status == "DATA_ERROR"
    assert result.event is None
    assert p == before


def test_cursor_ignores_other_trades_and_pre_entry_legacy_events():
    wb = book()
    ws = wb["Position Events"]
    for recorded, triggered, identity in [
        ("2026-09-08T13:26:07Z", "2026-09-01T14:11:00Z", ""),
        ("2026-09-08T16:00:00Z", "2026-09-08T15:00:00Z", "old-trade"),
        ("2026-09-08T14:00:00Z", "2026-09-08T13:59:00Z", "new-trade"),
    ]:
        ws.append([recorded, "run", "CHTR", "TAKE_PARTIAL_PROFIT", triggered] + [None] * 10 + [identity])
    assert read_last_event_times(wb, {"CHTR": position()})["CHTR"].isoformat() == "2026-09-08T13:59:00+00:00"


@pytest.mark.parametrize("bad_time,bad_id", [(True, False), (False, True)])
def test_invalid_event_cannot_mutate_ledger(bad_time, bad_id):
    p = position()
    event = build_event(p, action="EXIT_STOP", triggered_at="2026-09-01T14:11:00Z" if bad_time else "2026-09-08T14:11:00Z",
                        trigger_price=146.49, high=147, low=146, close=146.5,
                        quantity=26, currency_rate=1, note="test")
    if bad_id:
        event.trade_id = "another-trade"
    wb = book()
    with pytest.raises(ValueError, match="current trade"):
        apply_event(wb, {"CHTR": p}, p, event, "2026-09-08T14:12:00Z", "run", 1)
    assert wb["Trade Log"].max_row == 1
    assert wb["Position Events"].max_row == 1
    assert p["quantity"] == 26


def test_tp1_is_bound_to_trade_and_advances_stop_and_cursor():
    p = position()
    event = build_event(p, action="TAKE_PARTIAL_PROFIT", triggered_at="2026-09-08T14:11:00Z",
                        trigger_price=159.44, high=160, low=158, close=159.5,
                        quantity=13, currency_rate=1, note="test")
    wb = book()
    positions = {"CHTR": p}
    apply_event(wb, positions, p, event, "2026-09-08T14:12:00Z", "run", 1)
    assert p["quantity"] == 13
    assert p["stop_loss"] == p["entry_price"]
    assert p["risk_ils"] == 0
    assert p["partial_taken"] is True
    assert wb["Position Events"].cell(2, 16).value == "new-trade"
    assert read_last_event_times(wb, positions)["CHTR"].isoformat() == "2026-09-08T14:11:00+00:00"


def test_legacy_identity_includes_entry_timestamp():
    p = position()
    p["decision_json"] = ""
    first = position_trade_id(p)
    p["entry_date"] = "2026-09-08T14:00:00Z"
    assert first != position_trade_id(p)
