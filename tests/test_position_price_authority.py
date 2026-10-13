from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from app import strategy, position_market_data as data
from agent import market_lens_ui_agent as scanner, position_monitor as monitor


def frame(values, times):
    return pd.DataFrame(values, columns=['Open', 'High', 'Low', 'Close'], index=pd.to_datetime(times, utc=True))


def settings(tmp_path):
    return monitor.MonitorSettings(tmp_path / 'unused.xlsx', tmp_path, '5d', '1m', False, 'https://example.invalid')


def position():
    return dict(ticker='TEST', entry_date='2026-10-08T18:00:00Z', entry_price=100., current_price=101.,
                quantity=10, stop_loss=95., target_1=110., target_2=120., exposure_ils=1010.,
                risk_ils=50., partial_taken=False, notes='', decision_json='{"trade_id":"test-lot"}')


@pytest.mark.parametrize('ticker,setup_price,mark_price,stop', [
    ('TRGP', 283.96, 288.49, 286.5682), ('TGTX', 53.34, 55.05, 53.6524), ('PFE', 27.82, 28.27, 27.9217),
    ('TEST', 125., 100., 95.),
])
def test_setup_price_can_never_close_a_position(monkeypatch, ticker, setup_price, mark_price, stop):
    monkeypatch.setattr(strategy, 'fetch_position_mark', lambda _: dict(position_mark_price=mark_price, position_market_data_status='VALIDATED'))
    p = position()
    p.update(ticker=ticker, stop_loss=stop)
    before = deepcopy(p)
    result = scanner.SetupResult(ticker, 'No Trade', 0, setup_price, None, None, None, None, None, 0, '', '')
    decision = scanner.decide(result, open_positions={ticker:p}, cash=90000, exposure=1000, usd_ils=1,
                               max_position=10000, max_total_exposure=40000, max_risk=1000, min_rr=2,
                               sector_map={}, sector_health={})
    assert decision.action == 'HOLD' and decision.quantity == 0 and decision.cash_in_ils == 0
    assert decision.decision_json['position_mark_price'] == mark_price and p == before


def test_missing_intraday_data_preserves_mark_and_cannot_fake_an_exit(monkeypatch):
    monkeypatch.setattr(strategy, 'fetch_position_mark', lambda _: (_ for _ in ()).throw(ValueError('stale provider')))
    p = position()
    result = scanner.SetupResult('TEST', 'No Trade', 0, 1., None, None, None, None, None, 0, '', '')
    d = scanner.decide(result, open_positions={'TEST':p}, cash=90000, exposure=1010, usd_ils=1,
                       max_position=10000, max_total_exposure=40000, max_risk=1000, min_rr=2,
                       sector_map={}, sector_health={})
    assert d.action == 'HOLD' and d.decision_json['position_market_data_status'] == 'DATA_ERROR'
    assert scanner.refresh_open_position(p, result, 1, d) == 0
    assert p['current_price'] == 101 and p['quantity'] == 10 and p['stop_loss'] == 95


def test_live_or_future_bars_do_not_trigger_a_stop(monkeypatch, tmp_path):
    bars = frame([[101,102,100,101], [101,102,90,91]], ['2026-10-08T19:59:00Z','2026-10-08T20:00:00Z'])
    monkeypatch.setattr(monitor, 'fetch_intraday_frame', lambda *a, **kw: bars)
    result = monitor.monitor_position(position(), settings=settings(tmp_path), since=None, currency_rate=1,
                                       as_of=datetime(2026,10,8,20,0,30,tzinfo=timezone.utc))
    assert result.event is None and result.current_price == 101


def test_stale_regular_session_feed_is_data_error_not_healthy_hold(tmp_path):
    bars = frame([[101,102,100,101]], ['2026-10-07T19:59:00Z'])
    result = monitor.monitor_position(position(),settings=settings(tmp_path),since=None,currency_rate=1,
               frame_cache={'TEST':bars},as_of=datetime(2026,10,8,19,0,tzinfo=timezone.utc))
    assert result.status == 'DATA_ERROR' and result.event is None and result.current_price == 0


@pytest.mark.parametrize('terminal,second', [('EXIT_STOP',[101,102,99,100]), ('TAKE_PROFIT',[120,122,119,121])])
def test_real_completed_tp1_moves_stop_and_next_bar_executes_once(monkeypatch,tmp_path,terminal,second):
    bars = frame([[109,111,108,110],second], ['2026-10-08T18:01:00Z','2026-10-08T18:02:00Z'])
    monkeypatch.setattr(monitor,'fetch_intraday_frame',lambda *a, **kw:bars)
    wb = Workbook()
    wb.active.title='Trade Log'
    wb.active.append(['header'] * 32)
    monitor.ensure_position_events_sheet(wb)
    p=position()
    positions={'TEST':p}
    results,notifications=monitor.process_position_window(wb,positions,p,settings=settings(tmp_path),since=None,
               currency_rate=1,timestamp='2026-10-08T18:03:00Z',run_id='isolated')
    assert [r.event.action for r in results] == ['TAKE_PARTIAL_PROFIT',terminal]
    assert notifications[1][0]['stop_loss'] == 100 and notifications[1][0]['quantity'] == 5
    assert positions == {} and wb['Trade Log'].max_row == 3 and wb['Position Events'].max_row == 3
    # Persisted event cursor makes a restart skip the already sold partial.
    p=position()
    p.update(quantity=5, stop_loss=100, partial_taken=True)
    result=monitor.monitor_position(p,settings=settings(tmp_path),since=monitor.parse_timestamp('2026-10-08T18:01:00Z'),
                  currency_rate=1,frame_cache={'TEST':bars},as_of=datetime(2026,10,8,18,3,tzinfo=timezone.utc))
    assert result.event.action == terminal and result.event.quantity == 5


def test_extended_stops_remain_disabled_and_targets_keep_confirmation_policy(tmp_path):
    p=position()
    bars=frame([[101,102,90,91], [109,111,109,110], [110,112,109,111]],
                ['2026-10-08T20:10:00Z','2026-10-08T20:11:00Z','2026-10-08T20:12:00Z'])
    result=monitor.monitor_position(p,settings=settings(tmp_path),since=None,currency_rate=1,frame_cache={'TEST':bars},
              as_of=datetime(2026,10,8,20,13,tzinfo=timezone.utc))
    assert result.event.action == 'TAKE_PARTIAL_PROFIT' and result.event.triggered_at.startswith('2026-10-08T20:12')


def test_position_mark_is_from_completed_intraday_not_daily(monkeypatch):
    bars=frame([[101,102,100,101], [101,120,90,115]], ['2026-10-08T19:00:00Z','2026-10-08T19:01:00Z'])
    monkeypatch.setattr(data,'fetch_intraday_frame',lambda *a,**kw:bars)
    mark=data.fetch_position_mark('TEST',now=datetime(2026,10,8,19,1,30,tzinfo=timezone.utc))
    assert mark['position_mark_price'] == 101 and mark['position_market_data_status'] == 'VALIDATED'
    assert mark['position_mark_timestamp'].startswith('2026-10-08T19:00')


def test_unfinished_inconsistent_provider_bar_does_not_poison_completed_history():
    bars=frame([[101,102,100,101], [101,95,105,115]], ['2026-10-08T19:00:00Z','2026-10-08T19:01:00Z'])
    valid=data.validated_position_frame(bars,interval='1m',now=datetime(2026,10,8,19,1,30,tzinfo=timezone.utc))
    assert len(valid)==1 and valid['Close'].iloc[-1]==101
