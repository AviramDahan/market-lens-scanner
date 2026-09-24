from __future__ import annotations

import io
import json
import os
import re
import time
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook


@dataclass(frozen=True)
class SmokeCheck:
    name: str
    ok: bool
    detail: str
    duration_seconds: float


class SmokeFailure(RuntimeError):
    pass


class ProductionSmoke:
    def __init__(self) -> None:
        self.base_url = os.getenv(
            "MARKET_LENS_SMOKE_URL",
            "https://market-lens-scanner-fb63.onrender.com",
        ).rstrip("/")
        self.ticker = os.getenv("MARKET_LENS_SMOKE_TICKER", "MSFT").strip().upper()
        self.expected_revision = os.getenv("MARKET_LENS_SMOKE_EXPECTED_REVISION", "").strip()
        self.deploy_wait_seconds = env_int("MARKET_LENS_SMOKE_DEPLOY_WAIT_SECONDS", 360)
        self.monitor_max_age_minutes = env_int("MARKET_LENS_SMOKE_MONITOR_MAX_AGE_MINUTES", 30)
        self.report_path = Path(
            os.getenv("MARKET_LENS_SMOKE_REPORT_PATH", "production-smoke-report.json")
        )
        self.checks: list[SmokeCheck] = []
        self.dashboard: dict[str, Any] = {}

    def run(self) -> None:
        self.check("Render deployment", self.wait_for_deployment)
        self.check("Agent HTML", self.check_agent_html)
        self.check("Dashboard contract", self.check_dashboard)
        self.check("Decision JSONL", self.check_decision_jsonl)
        self.check("Tracker workbook", self.check_tracker)
        self.check("Read-only ticker scan", self.check_read_only_scan)
        self.check("Position monitor state", self.check_monitor_state)
        self.check("Render shadow monitor", self.check_render_shadow_monitor)

    def check(self, name: str, function) -> None:
        started = time.monotonic()
        try:
            detail = str(function())
        except Exception as exc:
            duration = round(time.monotonic() - started, 3)
            detail = str(exc)[:1000]
            self.checks.append(SmokeCheck(name, False, detail, duration))
            print(f"[FAIL] {name}: {detail}")
            raise SmokeFailure(f"{name}: {detail}") from exc
        duration = round(time.monotonic() - started, 3)
        self.checks.append(SmokeCheck(name, True, detail, duration))
        print(f"[OK] {name}: {detail}")

    def wait_for_deployment(self) -> str:
        deadline = time.monotonic() + self.deploy_wait_seconds
        last_detail = "No health response received."
        while True:
            try:
                payload = self.request_json("/health", attempts=1, timeout=30)
                validate_health(payload, self.expected_revision)
                revision = str(payload.get("revision") or "unknown")
                return f"status=ok; revision={revision[:12]}"
            except Exception as exc:
                last_detail = str(exc)
                if time.monotonic() >= deadline:
                    raise SmokeFailure(
                        f"Render did not expose the expected healthy revision within "
                        f"{self.deploy_wait_seconds}s: {last_detail}"
                    ) from exc
                time.sleep(10)

    def check_agent_html(self) -> str:
        body, headers = self.request_bytes("/agent", attempts=3, timeout=30)
        text = body.decode("utf-8", errors="replace")
        if "Agent Dashboard" not in text or len(text) < 1000:
            raise SmokeFailure("Agent dashboard HTML is missing or incomplete.")
        return f"bytes={len(body)}; content_type={headers.get('content-type', 'unknown')}"

    def check_dashboard(self) -> str:
        self.dashboard = self.request_json("/agent/data", attempts=3, timeout=60)
        return validate_dashboard(self.dashboard)

    def check_decision_jsonl(self) -> str:
        latest_run = self.dashboard.get("latest_run") or {}
        decision_url = decision_jsonl_url(latest_run)
        body, _headers = self.request_bytes(decision_url, attempts=3, timeout=60)
        count = validate_decision_jsonl(body)
        return f"records={count}; url={decision_url}"

    def check_tracker(self) -> str:
        body, _headers = self.request_bytes("/agent/tracker", attempts=3, timeout=90)
        sheets = validate_tracker(body)
        return f"bytes={len(body)}; sheets={len(sheets)}"

    def check_read_only_scan(self) -> str:
        payload = {
            "tickers": [self.ticker],
            "min_rr": 2.0,
            "analysis_period": "6mo",
            "session_id": "production-smoke-read-only",
            "include_charts": True,
            "persist_setups": False,
        }
        response = self.request_json(
            "/ui/scan",
            method="POST",
            payload=payload,
            attempts=2,
            timeout=150,
        )
        chart_url = validate_read_only_scan(response, self.ticker)
        chart, headers = self.request_bytes(chart_url, attempts=3, timeout=60)
        validate_chart(chart, headers)
        return f"ticker={self.ticker}; chart_bytes={len(chart)}; persisted=0"

    def check_monitor_state(self) -> str:
        detail = validate_monitor_state(
            self.dashboard,
            max_age_minutes=self.monitor_max_age_minutes,
        )
        return detail

    def check_render_shadow_monitor(self) -> str:
        payload = self.request_json(
            "/agent/monitor-shadow-status",
            attempts=3,
            timeout=30,
        )
        return validate_render_shadow_monitor(payload)

    def request_json(
        self,
        path: str,
        *,
        method: str = "GET",
        payload: dict[str, Any] | None = None,
        attempts: int = 1,
        timeout: int = 30,
    ) -> dict[str, Any]:
        body, _headers = self.request_bytes(
            path,
            method=method,
            payload=payload,
            attempts=attempts,
            timeout=timeout,
        )
        try:
            decoded = json.loads(body.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise SmokeFailure(f"{path} did not return valid JSON.") from exc
        if not isinstance(decoded, dict):
            raise SmokeFailure(f"{path} returned a non-object JSON payload.")
        return decoded

    def request_bytes(
        self,
        path: str,
        *,
        method: str = "GET",
        payload: dict[str, Any] | None = None,
        attempts: int = 1,
        timeout: int = 30,
    ) -> tuple[bytes, dict[str, str]]:
        url = urljoin(f"{self.base_url}/", path.lstrip("/"))
        encoded = json.dumps(payload).encode("utf-8") if payload is not None else None
        headers = {"User-Agent": "market-lens-production-smoke/1.0"}
        if encoded is not None:
            headers["Content-Type"] = "application/json"
        last_error: Exception | None = None
        for attempt in range(1, max(1, attempts) + 1):
            try:
                request = Request(url, data=encoded, headers=headers, method=method)
                with urlopen(request, timeout=timeout) as response:
                    status = int(getattr(response, "status", 0) or response.getcode())
                    body = response.read()
                    response_headers = {
                        key.lower(): value for key, value in response.headers.items()
                    }
                if not 200 <= status < 300:
                    raise SmokeFailure(f"{path} returned HTTP {status}.")
                return body, response_headers
            except (HTTPError, URLError, TimeoutError, SmokeFailure) as exc:
                last_error = exc
                if attempt < attempts:
                    time.sleep(3 * attempt)
        raise SmokeFailure(f"{method} {path} failed after {attempts} attempt(s): {last_error}")

    def write_report(self, *, status: str, error: str = "") -> None:
        payload = {
            "schema_version": 1,
            "status": status,
            "base_url": self.base_url,
            "ticker": self.ticker,
            "expected_revision": self.expected_revision,
            "error": error,
            "checks": [asdict(check) for check in self.checks],
        }
        self.report_path.write_text(
            json.dumps(payload, indent=2, sort_keys=True),
            encoding="utf-8",
        )


def validate_health(payload: dict[str, Any], expected_revision: str = "") -> None:
    if payload.get("status") != "ok":
        raise SmokeFailure(f"Health status is {payload.get('status')!r}, expected 'ok'.")
    if expected_revision:
        actual = str(payload.get("revision") or "").strip()
        if not actual or not revisions_match(actual, expected_revision):
            raise SmokeFailure(
                f"Render revision {actual[:12] or 'missing'} does not match expected "
                f"{expected_revision[:12]}."
            )


def validate_render_shadow_monitor(payload: dict[str, Any]) -> str:
    if payload.get("mode") != "shadow":
        raise SmokeFailure("Render monitor is not in shadow mode.")
    if payload.get("side_effects_enabled") is not False:
        raise SmokeFailure("Render shadow monitor unexpectedly exposes side effects.")
    if payload.get("enabled") is not True or payload.get("running") is not True:
        raise SmokeFailure("Render shadow monitor is not enabled and running.")
    if not isinstance(payload.get("event_journal"), list):
        raise SmokeFailure("Render shadow event journal is missing.")
    if int(payload.get("event_journal_limit") or 0) <= 0:
        raise SmokeFailure("Render shadow event journal limit is invalid.")
    return (
        f"status={payload.get('status')}; polls={int(payload.get('total_polls') or 0)}; "
        f"unique_events={int(payload.get('total_unique_events') or 0)}; "
        "side_effects=false"
    )


def revisions_match(actual: str, expected: str) -> bool:
    actual = actual.strip().lower()
    expected = expected.strip().lower()
    return bool(
        actual and expected and (actual.startswith(expected) or expected.startswith(actual))
    )


def validate_dashboard(payload: dict[str, Any]) -> str:
    if payload.get("status") != "ok":
        raise SmokeFailure(f"Dashboard status is {payload.get('status')!r}.")
    latest = payload.get("latest_run")
    if not isinstance(latest, dict):
        raise SmokeFailure("Dashboard latest_run is missing.")
    run_status = str(latest.get("run_status") or "")
    if run_status not in {"COMPLETE", "PARTIAL_OK", "OK", "SUCCESS"}:
        raise SmokeFailure(f"Latest run status is not successful: {run_status or 'missing'}.")
    coverage = latest.get("scan_coverage") if isinstance(latest.get("scan_coverage"), dict) else {}
    received = int(coverage.get("received") or len(latest.get("tickers") or []))
    if received <= 0:
        raise SmokeFailure("Dashboard reports zero result cards.")
    summary = str(latest.get("summary_text") or "")
    if "AUTH_FAILED" in summary or "RUN_FAILED" in summary:
        raise SmokeFailure("Latest dashboard summary contains a fatal run marker.")
    return f"run_id={latest.get('run_id')}; status={run_status}; results={received}"


def validate_decision_jsonl(body: bytes) -> int:
    count = 0
    tickers: set[str] = set()
    for line_number, raw_line in enumerate(body.decode("utf-8").splitlines(), start=1):
        if not raw_line.strip():
            continue
        try:
            record = json.loads(raw_line)
        except json.JSONDecodeError as exc:
            raise SmokeFailure(f"Decision JSONL line {line_number} is invalid.") from exc
        ticker = str(record.get("ticker") or "").strip().upper()
        if not ticker or not record.get("final_action") or not record.get("reason"):
            raise SmokeFailure(f"Decision JSONL line {line_number} is missing required fields.")
        if ticker in tickers:
            raise SmokeFailure(f"Decision JSONL contains duplicate ticker {ticker}.")
        tickers.add(ticker)
        count += 1
    if count == 0:
        raise SmokeFailure("Decision JSONL contains no records.")
    return count


def decision_jsonl_url(latest_run: dict[str, Any]) -> str:
    explicit = str(latest_run.get("decision_jsonl_url") or "")
    if explicit.startswith("/agent-results/decisions/") and ".." not in explicit:
        return explicit
    run_id = str(latest_run.get("run_id") or "")
    if re.fullmatch(r"\d{8}_\d{6}", run_id):
        return f"/agent-results/decisions/market_lens_agent_{run_id}.jsonl"
    raise SmokeFailure("Latest run does not expose a safe Decision JSONL URL or run id.")


def validate_tracker(body: bytes) -> list[str]:
    if not body.startswith(b"PK"):
        raise SmokeFailure("Tracker response is not an XLSX zip payload.")
    try:
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise SmokeFailure(
            f"Tracker workbook could not be opened: {exc.__class__.__name__}."
        ) from exc
    try:
        sheets = list(workbook.sheetnames)
    finally:
        workbook.close()
    required = {"Dashboard", "Trade Log", "Open Positions", "Position Events"}
    missing = sorted(required - set(sheets))
    if missing:
        raise SmokeFailure(f"Tracker is missing required sheets: {', '.join(missing)}.")
    return sheets


def validate_read_only_scan(payload: dict[str, Any], ticker: str) -> str:
    results = payload.get("results") if isinstance(payload.get("results"), list) else []
    errors = payload.get("errors") if isinstance(payload.get("errors"), dict) else {}
    saved = payload.get("saved_setups") if isinstance(payload.get("saved_setups"), list) else []
    result = next(
        (item for item in results if str(item.get("ticker") or "").upper() == ticker),
        None,
    )
    if result is None:
        raise SmokeFailure(f"Read-only scan returned no {ticker} result; errors={errors}.")
    if saved:
        raise SmokeFailure("Read-only scan unexpectedly persisted setup records.")
    charts = payload.get("charts") if isinstance(payload.get("charts"), dict) else {}
    chart_url = str(charts.get(ticker) or "")
    if not chart_url.startswith("/charts/"):
        raise SmokeFailure(f"Read-only scan did not return a chart for {ticker}.")
    return chart_url


def validate_chart(body: bytes, headers: dict[str, str]) -> None:
    if len(body) < 1_000 or not body.startswith(b"\x89PNG\r\n\x1a\n"):
        raise SmokeFailure("Generated chart is not a valid non-empty PNG.")
    content_type = headers.get("content-type", "")
    if "image/png" not in content_type:
        raise SmokeFailure(f"Generated chart content type is {content_type or 'missing'}.")


def validate_monitor_state(payload: dict[str, Any], *, max_age_minutes: int) -> str:
    health = payload.get("system_health")
    if not isinstance(health, dict):
        raise SmokeFailure("Dashboard system_health is missing.")
    status = str(health.get("latest_monitor_status") or "")
    if status != "MONITOR_OK":
        raise SmokeFailure(f"Latest position monitor status is {status or 'missing'}.")
    failures = int(health.get("latest_monitor_positions_failed") or 0)
    if failures:
        raise SmokeFailure(f"Latest position monitor has {failures} failed position(s).")
    age = health.get("latest_monitor_age_minutes")
    if max_age_minutes > 0 and age is not None and int(age) > max_age_minutes:
        raise SmokeFailure(
            f"Latest position monitor is stale: {age} minutes (maximum {max_age_minutes})."
        )
    events = int(health.get("latest_monitor_event_count") or 0)
    checked = int(health.get("latest_monitor_positions_checked") or 0)
    return f"status={status}; checked={checked}; events={events}; dispatch_attempted=false"


def env_int(name: str, default: int) -> int:
    try:
        return int(os.getenv(name, str(default)))
    except ValueError:
        return default


def main() -> None:
    smoke = ProductionSmoke()
    try:
        smoke.run()
    except Exception as exc:
        smoke.write_report(status="failed", error=str(exc))
        raise SystemExit(1) from exc
    smoke.write_report(status="passed")
    print(f"Production smoke passed; report={smoke.report_path}")


if __name__ == "__main__":
    main()
