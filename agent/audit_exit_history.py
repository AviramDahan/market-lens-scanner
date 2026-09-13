"""Read-only audit of recorded monitor evidence; never repairs cash or trades."""

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from agent.position_monitor import parse_timestamp
from app.agent_dashboard import compute_full_trade_performance, compute_realized_pnl, read_trades


def audit_workbook(wb):
    findings = []
    trade_records = read_trades(wb)
    trades = list(wb["Trade Log"].iter_rows(min_row=2, values_only=True))
    buys = {}
    exits = {}
    for row in trades:
        row = tuple(row) + (None,) * max(0, 21 - len(row))
        identity = str(row[20] or "")
        if row[1] == "BUY_SIMULATED" and identity:
            buys.setdefault(identity, []).append(row)
        elif row[1] in {"EXIT_STOP", "TAKE_PROFIT", "TAKE_PARTIAL_PROFIT"}:
            key = (str(row[2]), str(row[1]), parse_timestamp(row[0]))
            exits.setdefault(key, []).append(row)
    events = list(wb["Position Events"].iter_rows(min_row=2, values_only=True)) if "Position Events" in wb.sheetnames else []
    checked = 0
    matched_events = {}
    for number, raw in enumerate(events, 2):
        row = tuple(raw) + (None,) * max(0, 16 - len(raw))
        if not row[2] or str(row[3] or "").startswith("VOID"):
            continue
        checked += 1
        recorded, triggered = parse_timestamp(row[0]), parse_timestamp(row[4])
        matches = exits.get((str(row[2]), str(row[3]), recorded), [])
        item = {"event_row": number, "ticker": row[2], "action": row[3],
                "recorded_at": str(row[0]), "triggered_at": str(row[4])}
        if len(matches) != 1:
            findings.append({**item, "status": "UNRESOLVED_EXIT_IDENTITY"})
            continue
        trade = matches[0]
        identity = str(trade[20] or "")
        item["trade_id"] = identity
        if identity:
            matched_events[identity] = matched_events.get(identity, 0) + 1
        entries = buys.get(identity, [])
        if len(entries) != 1 or str(entries[0][2]) != str(row[2]):
            findings.append({**item, "status": "UNRESOLVED_ENTRY_IDENTITY"})
            continue
        entry = parse_timestamp(entries[0][0])
        item["entry_at"] = str(entries[0][0])
        if row[15] and str(row[15]) != identity:
            status = "TRADE_ID_MISMATCH"
        elif entry is None or recorded is None or triggered is None:
            status = "MISSING_TIMESTAMP"
        elif triggered < entry:
            status = "PRE_ENTRY_EVENT"
        elif triggered > recorded:
            status = "FUTURE_EVENT"
        else:
            continue
        findings.append({**item, "status": status})
    return {"audit_version": "exit_identity_v1", "read_only": True,
            "events_checked": checked, "findings": findings,
            "measurement_provenance": measurement_provenance(buys, matched_events, findings),
            "accounting_reconciliation": accounting_reconciliation(trade_records),
            "affected_trade_ids": sorted({f["trade_id"] for f in findings if f.get("trade_id")}),
            "limitations": ["Absence of findings does not validate fills, prices, or data completeness.",
                            "Only recorded monitor events are audited; scanner exits without events are not covered.",
                            "No corrected PnL is inferred from invalid execution evidence."]}


def accounting_reconciliation(trades):
    """Explain realized-PnL scopes without changing or validating historical fills."""
    exit_events = compute_realized_pnl(trades)
    lifecycle = compute_full_trade_performance(trades)
    event_realized = round(float(exit_events["total"]), 2)
    closed_realized = round(float(lifecycle["closed_trade_realized_pnl_ils"]), 2)
    open_partial_realized = round(float(lifecycle["open_lot_partial_realized_pnl_ils"]), 2)
    lifecycle_realized = round(float(lifecycle["all_lifecycle_realized_pnl_ils"]), 2)
    return {
        "version": "realized_pnl_scope_v1",
        "currency_field_suffix": "legacy_ils_name; workbook currency setting controls display currency",
        "exit_event_realized_pnl_ils": event_realized,
        "closed_trade_realized_pnl_ils": closed_realized,
        "open_lot_partial_realized_pnl_ils": open_partial_realized,
        "all_lifecycle_realized_pnl_ils": lifecycle_realized,
        "reconciliation_delta_ils": round(event_realized - lifecycle_realized, 2),
        "reconciled": abs(event_realized - lifecycle_realized) < 0.01,
        "closed_trade_count": lifecycle["closed_count"],
        "open_trade_count": lifecycle["open_count"],
        "cash_or_pnl_corrected": False,
        "usage": (
            "Closed-trade PnL excludes partial exits from positions that remain open; "
            "all-lifecycle realized PnL includes both scopes."
        ),
    }


def measurement_provenance(buys, matched_events, findings):
    """Annotate evidence quality, without presenting a filtered portfolio as a backtest."""
    trades = []
    for identity in sorted(set(buys) | set(matched_events)):
        issues = [f for f in findings if f.get("trade_id") == identity]
        count = matched_events.get(identity, 0)
        trades.append({
            "trade_id": identity,
            "status": "NEEDS_RECONCILIATION" if issues else (
                "IDENTITY_CHECK_ONLY" if count else "NO_MATCHED_MONITOR_EVIDENCE"
            ),
            "matched_monitor_events": count,
            "finding_statuses": sorted({f["status"] for f in issues}),
            "fill_prices_validated": False,
        })
    return {
        "version": "historical_evidence_v1",
        "trades": trades,
        "unattributed_findings": sum(not f.get("trade_id") for f in findings),
        "cash_or_pnl_corrected": False,
        "downstream_capital_effects_replayed": False,
        "usage": "Do not treat no finding as a verified fill or exclude trades to claim improved returns.",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.resolve() == args.workbook.resolve():
        parser.error("Output must not overwrite the input workbook")
    digest = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        report = audit_workbook(wb)
    finally:
        wb.close()
    report["workbook_sha256"] = digest
    if hashlib.sha256(args.workbook.read_bytes()).hexdigest() != digest:
        raise RuntimeError("Workbook changed during audit; discard this result and retry")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({"events_checked": report["events_checked"], "findings": len(report["findings"]),
                      "output": str(args.output)}))


if __name__ == "__main__":
    main()
